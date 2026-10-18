"""ds-chat —— 终端里的 AI 聊天 + Agent，界面仿 Grok 风格。

可以保存多套 API 配置（如 DeepSeek 中转站、Claude 官方），聊天时用 /api 切换。

Agent 模式下 AI 可以运行命令、读写文件；运行命令和改文件前都会先问用户。

配置保存在 ~/.config/ds-chat/config.json，第一次运行时会引导填写。
"""
import asyncio
import base64
import gzip
import html
import json
import os
import re
import select
import selectors
import shutil
import signal
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from contextlib import contextmanager

from prompt_toolkit import PromptSession
from prompt_toolkit.cursor_shapes import CursorShape
from prompt_toolkit.formatted_text import HTML
from prompt_toolkit.history import FileHistory
from prompt_toolkit.styles import Style
from prompt_toolkit.utils import get_cwidth
from rich.console import Console, Group
from rich.live import Live
from rich.markdown import Markdown
from rich.padding import Padding
from rich.panel import Panel
from rich.segment import Segment
from rich.spinner import Spinner
from rich.table import Table
from rich.text import Text

VERSION = "1.3.4"
# 有的中转站会拦截 Python 默认的 User-Agent（返回 403），所以带上自己的名字
USER_AGENT = f"ds-chat/{VERSION}"
INSTALL_CMD = "curl -fsSL https://raw.githubusercontent.com/imthnio/ds-chat/main/install.sh | bash"
APP_DIR = os.path.expanduser("~/.local/share/ds-chat")
CONFIG_DIR = os.path.expanduser("~/.config/ds-chat")
CONFIG_FILE = os.path.join(CONFIG_DIR, "config.json")

# 安装向导里可以直接选的线路
PRESET_LINES = [
    ("主线路", "https://ai.clol.site/v1"),
    ("备用线路", "https://ai.bcol.cc/v1"),
]

ACCENT = "#7aa2ff"  # 主题色，随当前模型变化（见 apply_theme）
DIM = "grey50"

# 顶部 logo：按桌面上的图片取色画成的像素图，每个字符是一个像素，空格是透明。
# 画的时候上下两行像素合成一个字符（▀ 上半块用前景色、下半块用背景色）。
LOGOS = {
    # Claude：橘色小螃蟹（Claw'd），原图是 12x8 的像素画，这里放大 2 倍
    'claude': ({'o': '#da7758', 'k': '#000000'}, [
        '    oooooooooooooooo    ',
        '    oooooooooooooooo    ',
        '    ookkooooooookkoo    ',
        '    ookkooooooookkoo    ',
        'oooooooooooooooooooooooo',
        'oooooooooooooooooooooooo',
        'oooooooooooooooooooooooo',
        'oooooooooooooooooooooooo',
        '    oooooooooooooooo    ',
        '    oooooooooooooooo    ',
        '    oooooooooooooooo    ',
        '    oooooooooooooooo    ',
        '    oo  oo    oo  oo    ',
        '    oo  oo    oo  oo    ',
        '    oo  oo    oo  oo    ',
        '    oo  oo    oo  oo    ',
    ]),
    # DeepSeek：蓝色鲸鱼
    'deepseek': ({'d': '#4d6bfe', 'a': '#ffffff'}, [
        '                    d         ',
        '       dddddddd     dd        ',
        '    ddddddddddd     ddd      d',
        '   ddddddddddddd    dddd  dddd',
        '  ddddddddddddddd   dddddddddd',
        ' ddddddddddddddddd   dddddddd ',
        ' dddddddddddddddddd   dddddd  ',
        'dddddddddddddddddddd ddddd    ',
        'ddaaaaadddddddddddddddddd     ',
        'ddaaaaaaadddddddddddddddd     ',
        'ddaaaaaaaadddddddaadddddd     ',
        'ddaaaaaaaaaddddddaadddddd     ',
        'dddaaaaaaaaadddddddddddd      ',
        'dddaaaaaaaaaaddddddddddd      ',
        ' dddaaaaaaaaaaddddddddd       ',
        ' dddaaaaaaaaaaddddddddd       ',
        '  dddaaaaadaaaaddddddd        ',
        '   dddaaaadddaaaddddd         ',
        '    ddddaaddddaaadddddd       ',
        '     ddddddddddddddddddd      ',
        '      dddddddddddd            ',
        '        ddddddd               ',
    ]),
    # Codex：紫到蓝渐变的云，中间白色 >_
    'codex': ({'a': '#c4b5fd', 'b': '#b1a6fe', 'c': '#9d98fe', 'd': '#8a8afe', 'e': '#767cfe', 'f': '#636efe', 'g': '#5462fe', 'w': '#ffffff'}, [
        '         aaaaa            ',
        '       aaaaaaaa           ',
        '      aaaaaaaaaaaaaaa     ',
        '     aaaaaaaaaaaaaaaaa    ',
        '     bbbbbbbbbbbbbbbbbb   ',
        '   bbbbbbbbbbbbbbbbbbbbb  ',
        '  bbbbbbbbbbbbbbbbbbbbbb  ',
        ' bbbbbbbbbbbbbbbbbbbbbbbb ',
        'ccccccwwccccccccccccccccc ',
        'cccccccwwccccccccccccccc  ',
        'ccccccccwwccccccccccccccc ',
        'cccccccccwwcccccccccccccc ',
        'dddddddddwwddddddddddddddd',
        'ddddddddwwdddddddddddddddd',
        ' ddddddwwddddddddddddddddd',
        ' eeeeewweeeeewwwwwwweeeeee',
        '  eeeeeeeeeeewwwwwwweeeeee',
        ' eeeeeeeeeeeeeeeeeeeeeeeee',
        ' eeeeeeeeeeeeeeeeeeeeeeee ',
        '  ffffffffffffffffffffff  ',
        '  fffffffffffffffffffff   ',
        '   ffffffffffffffffff     ',
        '    fffffffffffffffff     ',
        '     ggggggggggggggg      ',
        '          ggggggggg       ',
        '            gggggg        ',
    ]),
    # 其他模型：四角星
    'other': ({'a': '#c4b5fd', 'b': '#a5a0fe', 'c': '#8a8afe', 'd': '#7a90ff', 'e': '#7aa2ff'}, [
        '      a      ',
        '      a      ',
        '     bbb     ',
        '     bbb     ',
        '    ccccc    ',
        '  ddddddddd  ',
        'eeeeeeeeeeeee',
        '  ddddddddd  ',
        '    ccccc    ',
        '     bbb     ',
        '     bbb     ',
        '      a      ',
        '      a      ',
    ]),
}

if sys.platform == "darwin":
    # macOS 的 kqueue 不支持 /dev/tty（curl | bash 安装时就是从它读键盘），改用 select
    class _SelectLoopPolicy(asyncio.DefaultEventLoopPolicy):
        def new_event_loop(self):
            return asyncio.SelectorEventLoop(selectors.SelectSelector())

    asyncio.set_event_loop_policy(_SelectLoopPolicy())

console = Console(highlight=False)
_plain = None


def plain_prompt(*args, **kwargs):
    """菜单、向导用的简单输入（用到时才创建，免得 ai --version 之类也去接管终端）。"""
    global _plain
    if _plain is None:
        _plain = PromptSession()
    return _plain.prompt(*args, **kwargs)


# ───────────────────────── 配置 ─────────────────────────

# 配置文件里可以有多套 API 配置（profiles），active 是正在用的那套。
# 每套：name 名字、type（"openai" = OpenAI 格式，如 DeepSeek 中转站；"anthropic" = Claude 格式）、
# base_url、backup_url（可选）、api_key、model。
STORE = {"profiles": [], "active": 0, "agent": True}


def load_config():
    """读取配置，返回正在用的那套；没有配置时返回 None。旧版只有一套的配置会自动迁移。"""
    try:
        with open(CONFIG_FILE, encoding="utf-8") as f:
            data = json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        return None
    if not isinstance(data, dict):
        return None
    if "profiles" not in data and data.get("base_url"):
        data = {"profiles": [{"name": "DeepSeek", "type": "openai", **{
            k: data[k] for k in ("base_url", "backup_url", "api_key", "model") if k in data}}],
            "active": 0, "agent": data.get("agent", True)}
    profiles = [p for p in data.get("profiles") or []
                if isinstance(p, dict) and p.get("base_url") and p.get("api_key") and p.get("model")]
    if not profiles:
        return None
    STORE.update(profiles=profiles, agent=data.get("agent", True),
                 active=min(max(int(data.get("active") or 0), 0), len(profiles) - 1))
    return active_profile()


def active_profile():
    return STORE["profiles"][STORE["active"]]


def save_config(_cfg=None):
    """保存全部配置（传进来的 cfg 是 STORE 里某一套的引用，改它就是改 STORE）。"""
    os.makedirs(CONFIG_DIR, exist_ok=True)
    tmp = CONFIG_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(STORE, f, ensure_ascii=False, indent=2)
    os.chmod(tmp, 0o600)
    os.replace(tmp, CONFIG_FILE)


def kind_label(p):
    return "Claude 格式" if p.get("type") == "anthropic" else "OpenAI 格式"


def host(url):
    return re.sub(r"^https?://|/v1$", "", url or "")


def normalize_base(url):
    url = url.strip().rstrip("/")
    if not re.match(r"https?://", url):
        url = "https://" + url
    if not url.endswith("/v1"):
        url += "/v1"
    return url


# ── Claude（Anthropic Messages API，用官方 SDK）──

ANTHROPIC_OFFICIAL = "https://api.anthropic.com"
# 中转站列不出模型时可以选的 Claude 模型
KNOWN_CLAUDE_MODELS = ["claude-opus-5-5", "claude-sonnet-5-5", "claude-haiku-5-5", "claude-fable-5-1"]
# 这些模型支持服务端拒答兜底（被安全分类器拒绝时自动换模型重答），只在官方接口上开
FALLBACK_MODELS = {"claude-fable-5-1", "claude-opus-5-5", "claude-opus-5", "claude-sonnet-5-5"}
# 支持自适应思考、64K 输出的新模型；其他（老模型或中转站的非 Claude 模型）不传思考参数
_MODERN_CLAUDE = re.compile(r"claude-(opus-(4-[6-9]|5)|sonnet-(4-6|5)|fable|mythos|haiku-5)")
_anthropic_clients = {}


def is_official(p):
    return p.get("base_url", "").rstrip("/") == ANTHROPIC_OFFICIAL


def anthropic_client(p):
    import anthropic  # 只有用 Claude 时才加载，启动更快
    k = (p["base_url"], p["api_key"])
    if k not in _anthropic_clients:
        _anthropic_clients[k] = anthropic.Anthropic(api_key=p["api_key"], base_url=p["base_url"])
    return _anthropic_clients[k]


def to_anthropic_messages(history):
    """把内部（OpenAI 格式）的对话记录转成 Claude 的格式。Claude 自己产生的消息原样带回（含思考块）。"""
    out = []

    def add(role, blocks):
        if out and out[-1]["role"] == role:
            out[-1]["content"].extend(blocks)
        else:
            out.append({"role": role, "content": list(blocks)})

    def image(img):
        return {"type": "image", "source": {"type": "base64", "media_type": img["media_type"], "data": img["data"]}}

    for m in history:
        if m["role"] == "user":
            add("user", [{"type": "text", "text": m["content"]}] + [image(x) for x in m.get("images") or []])
        elif m["role"] == "assistant":
            if m.get("anthropic_content"):
                add("assistant", m["anthropic_content"])
                continue
            blocks = [{"type": "text", "text": m["content"]}] if m.get("content") else []
            for c in m.get("tool_calls") or []:
                try:
                    args = json.loads(c["function"].get("arguments") or "{}")
                except json.JSONDecodeError:
                    args = {}
                blocks.append({"type": "tool_use", "id": c["id"], "name": c["function"]["name"],
                               "input": args if isinstance(args, dict) else {}})
            if blocks:
                add("assistant", blocks)
        elif m["role"] == "tool":
            content = m["content"] or "（无输出）"
            if m.get("images"):
                content = [{"type": "text", "text": content}] + [image(x) for x in m["images"]]
            add("user", [{"type": "tool_result", "tool_use_id": m["tool_call_id"], "content": content}])
    return out


def stream_anthropic(p, system, history, tools=None):
    """和 stream_chat 产出同样的事件，最后多一个 ("final", {...})，带工具调用、原始内容块和结束原因。"""
    model = p["model"]
    official = is_official(p)
    kw = {"model": model, "system": system, "messages": to_anthropic_messages(history),
          "max_tokens": 64000 if _MODERN_CLAUDE.match(model) else 8192}
    if tools:
        kw["tools"] = [{"name": t["function"]["name"], "description": t["function"]["description"],
                        "input_schema": t["function"]["parameters"],
                        # 大段文件内容边生成边传；中转站可能不认这个字段，只在官方接口开
                        **({"eager_input_streaming": True} if official else {})} for t in tools]
    if _MODERN_CLAUDE.match(model):
        kw["thinking"] = {"type": "adaptive", "display": "summarized"}
    if official and model in FALLBACK_MODELS:
        kw["extra_headers"] = {"anthropic-beta": "server-side-fallback-2026-07-01"}
        kw["extra_body"] = {"fallbacks": "default"}
    with anthropic_client(p).messages.stream(**kw) as stream:
        for ev in stream:
            if ev.type == "content_block_start" and ev.content_block.type == "tool_use":
                yield "tool_pending", None
            elif ev.type == "content_block_delta":
                if ev.delta.type == "text_delta":
                    yield "content", ev.delta.text
                elif ev.delta.type == "thinking_delta":
                    yield "reasoning", ev.delta.thinking
        final = stream.get_final_message()
    yield "usage", {"total_tokens": final.usage.input_tokens + final.usage.output_tokens}
    calls = [{"id": b.id, "type": "function",
              "function": {"name": b.name, "arguments": json.dumps(b.input, ensure_ascii=False)}}
             for b in final.content if b.type == "tool_use"]
    details = getattr(final, "stop_details", None)
    yield "final", {"calls": calls, "raw": [b.to_dict() for b in final.content],
                    "stop_reason": final.stop_reason,
                    "refusal_category": getattr(details, "category", None) if details else None}


def list_models(p):
    if p.get("type") != "anthropic":
        return fetch_models(p["base_url"], p["api_key"])
    import anthropic
    try:
        return [m.id for m in anthropic_client(p).models.list()]
    except anthropic.NotFoundError:
        return list(KNOWN_CLAUDE_MODELS)  # 有的中转站没有模型列表接口


def error_code(e):
    """HTTP 状态码：urllib 的 HTTPError 用 .code，Claude SDK 的错误用 .status_code。"""
    return getattr(e, "status_code", None) or getattr(e, "code", None)


def fetch_models(base, key):
    req = urllib.request.Request(f"{base}/models", headers={"Authorization": f"Bearer {key}",
                                                            "User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=20) as resp:
        data = json.load(resp)
    return sorted(m["id"] for m in data.get("data", []) if m.get("id"))


def test_model(p, model):
    """用和聊天一样的流式请求试一次，收到第一块数据就算成功。"""
    if p.get("type") == "anthropic":
        with anthropic_client(p).messages.stream(
                model=model, max_tokens=64, messages=[{"role": "user", "content": "hi"}]) as stream:
            for _ in stream:
                return
        return
    for _ in stream_chat(p["base_url"], p["api_key"], model, [{"role": "user", "content": "hi"}]):
        return


def ask_choice(title, rows, default=1):
    """rows: [(名称, 说明)]，返回 1 开始的序号。"""
    table = Table(box=None, padding=(0, 2), show_header=False)
    for i, (name, desc) in enumerate(rows, 1):
        table.add_row(Text(f"{i}", style=f"bold {ACCENT}"), Text(name, style="bold"), Text(desc, style=DIM))
    console.print(Panel(table, title=title, title_align="left", border_style=ACCENT))
    while True:
        ans = plain_prompt(HTML(f"<b>输入序号</b> (直接回车 = {default}): ")).strip()
        if not ans:
            return default
        if ans.isdigit() and 1 <= int(ans) <= len(rows):
            return int(ans)
        console.print(Text("  请输入列表里的数字", style="red"))


# 向导第 1 步的选项：(名称, 说明, 类型, 地址)；地址为空表示让用户自己填
LINE_CHOICES = [
    ("DeepSeek 主线路", PRESET_LINES[0][1], "openai", PRESET_LINES[0][1]),
    ("DeepSeek 备用线路", PRESET_LINES[1][1], "openai", PRESET_LINES[1][1]),
    ("Claude 官方", ANTHROPIC_OFFICIAL, "anthropic", ANTHROPIC_OFFICIAL),
    ("其他中转站 · OpenAI 格式", "大多数中转站、DeepSeek、通义等", "openai", ""),
    ("其他中转站 · Claude 格式", "提供 Claude 接口的中转站", "anthropic", ""),
]


# 从 PDF、备忘录里复制时常混进来的“长得像”的字符
_LOOKALIKE = str.maketrans({"ɑ": "a", "ο": "o", "О": "O", "о": "o", "І": "I", "і": "i", "Ⅰ": "I", "ⅼ": "l",
                            "‐": "-", "‑": "-", "–": "-", "—": "-", "－": "-", "“": '"', "”": '"', "‘": "'", "’": "'"})


def parse_pasted(text):
    """清理粘贴进来的密钥或连接信息，返回 (密钥, 地址)，地址没有就是空字符串。
    支持中转站“复制连接信息”给的 JSON，如 {"key": "sk-…", "url": "https://…"}。"""
    text = text.translate(_LOOKALIKE).strip()
    url = ""
    m = re.search(r"\{.*\}", text, re.S)
    if m:
        try:
            # 复制时密钥、网址中间常被断行；这些字段里本来就没有空白，全部去掉再解析
            data = json.loads(re.sub(r"\s+", "", m.group(0)))
            text, url = str(data.get("key") or ""), str(data.get("url") or "")
        except (json.JSONDecodeError, AttributeError):
            pass
    key = re.sub(r"\s+", "", text).strip("\"'`")
    if re.fullmatch(r"sk[A-Za-z0-9]{48}", key):
        key = "sk-" + key[2:]  # 从 PDF 复制时，sk- 在行尾的话 - 会被当成断字符号吞掉
    return key, re.sub(r"\s+", "", url).strip("\"'`")


def setup_wizard(old=None):
    """填写一套 API 配置并测试，返回这套配置（不保存，由调用的地方保存）。old 是要修改的那套。"""
    console.clear()
    title = f"修改配置「{old['name']}」" if old else "添加一套 API 配置"
    console.print(Panel(Text.assemble(
        (f"{title}\n\n", f"bold {ACCENT}"),
        ("需要填：线路、API 密钥、默认模型，最后给它起个名字。\n", ""),
        ("随时按 Ctrl+C 可以取消。", DIM),
    ), border_style=ACCENT, padding=(1, 2)))
    old = old or {}

    while True:
        # 第 1 步：线路 / 类型
        default = next((i for i, c in enumerate(LINE_CHOICES, 1)
                        if c[3] and c[3] == old.get("base_url")), 1)
        if old and default == 1 and old.get("base_url") != LINE_CHOICES[0][3]:
            default = 5 if old.get("type") == "anthropic" else 4
        n = ask_choice("第 1 步 · 选择线路", [(c[0], c[1]) for c in LINE_CHOICES], default)
        name, _, kind, base = LINE_CHOICES[n - 1]
        if not base:
            url, pasted_key = "", ""
            while not url:
                url = plain_prompt(HTML("<b>接口地址</b> (例如 https://example.com，"
                                        "也可以直接粘贴中转站给的连接信息): ")).strip()
                pasted_key, from_json = parse_pasted(url)
                if from_json:  # 粘贴的是 {"key":…, "url":…}
                    url = from_json
                else:
                    url, pasted_key = url.translate(_LOOKALIKE).strip(" \"'"), ""
            base = normalize_base(url)
            if kind == "anthropic":
                base = re.sub(r"/v1$", "", base)  # Claude SDK 自己会加 /v1
        else:
            pasted_key = ""
        backup = next((u for _, u in PRESET_LINES if u != base), "") if base in dict(PRESET_LINES).values() else ""
        draft = {"type": kind, "base_url": base, "backup_url": backup}

        # 第 2 步：密钥
        console.print()
        same = old.get("api_key") and old.get("type") == kind
        hint = "（直接回车 = 沿用原来的密钥）" if same else ""
        console.print(Text(f"第 2 步 · 粘贴 API 密钥{hint}，输入时不会显示出来", style=f"bold {ACCENT}"))
        key = pasted_key
        if key:
            console.print(Text(f"  已从连接信息里读到密钥 {key[:6]}…{key[-4:]}", style="green"))
        while not key:
            raw = plain_prompt("🔑 ", is_password=True)
            key = parse_pasted(raw)[0] if raw.strip() else (old.get("api_key", "") if same else "")
            if key and key != raw.strip():
                console.print(Text(f"  已自动清理粘贴内容（去掉空格、换行、引号或长得像的字符）：{key[:6]}…{key[-4:]}",
                                   style="green"))
        if not key.isascii():
            console.print(Text("  注意：密钥里有特殊字符，可能是复制时混进来的，最好重新复制", style="yellow"))
        draft["api_key"] = key
        if kind == "openai" and key.startswith("sk-ant-"):
            console.print(Text("  提示：这像是 Claude 官方密钥，Claude 官方请在第 1 步选「Claude 官方」", style="yellow"))

        # 测试连接
        try:
            with console.status("正在测试连接…", spinner="dots"):
                models = list_models(draft)
        except Exception as e:
            code = error_code(e)
            detail = ""
            if hasattr(e, "read"):
                try:
                    detail = e.read().decode("utf-8", "replace").strip()[:200]
                except Exception:
                    pass
            elif getattr(e, "message", None):
                detail = str(e.message)[:200]
            if code == 401:
                msg = "密钥不对（401）。检查是不是复制完整了、有没有多余的字符"
            elif code == 403:
                msg = "服务器拒绝访问（403）：密钥没有权限，或者被中转站拦截了"
            elif code:
                msg = f"服务器返回错误 HTTP {code}"
            else:
                msg = ""
            if code and detail:
                msg += f"\n服务器说：{detail}"
            if not code:
                msg = f"连不上 {base}\n{e}\n换一条线路再试试"
            console.print(Panel(Text(msg + "\n请重新填写。"), title="连接失败", border_style="red"))
            continue
        if not models:
            console.print(Panel(Text("连接成功，但这个密钥没有可用的模型，请检查分组设置。"),
                                title="没有模型", border_style="red"))
            continue
        console.print(Text(f"  ✓ 连接成功，找到 {len(models)} 个模型\n", style="green"))
        break

    # 第 3 步：默认模型（选完实际试一次，有的中转站会列出用不了的模型）
    if kind == "anthropic":
        default = models.index("claude-opus-5-5") + 1 if "claude-opus-5-5" in models else 1
    else:
        default = next((i for i, m in enumerate(models, 1) if "pro" in m.lower() and "/" not in m), 1)
    if old.get("model") in models:
        default = models.index(old["model"]) + 1
    while True:
        n = ask_choice("第 3 步 · 选择默认模型", [(m, "") for m in models], default)
        model = models[n - 1]
        try:
            with console.status(f"正在测试 {model}…", spinner="dots"):
                test_model(draft, model)
            break
        except Exception as e:
            console.print(Panel(Text(f"{model} 现在用不了（{e}），请换一个。"),
                                title="模型不可用", border_style="red"))
            default = 1 if n != 1 else min(2, len(models))
    draft["model"] = model

    # 第 4 步：名字
    suggest = old.get("name") or ("Claude" if kind == "anthropic" else
                                  "DeepSeek" if "deepseek" in model.lower() else host(base))
    label = plain_prompt(HTML(f"\n<b>第 4 步 · 给这套配置起个名字</b> (直接回车 = {suggest}): ")).strip()
    draft["name"] = label[:20] or suggest

    console.print(Panel(Text.assemble(
        (f"✓ 配置「{draft['name']}」已就绪\n\n", "bold green"),
        ("聊天时输入 ", ""), ("/api", f"bold {ACCENT}"), (" 可以在几套配置之间切换。", ""),
    ), border_style="green", padding=(1, 2)))
    plain_prompt(HTML("<grey>按回车继续…</grey>"))
    return draft


def profile_menu(cfg):
    """/api：切换、添加、修改、删除 API 配置。返回切换后正在用的那套。"""
    while True:
        profiles = STORE["profiles"]
        table = Table(box=None, padding=(0, 2), show_header=False)
        for i, p in enumerate(profiles, 1):
            mark = "● 当前" if i - 1 == STORE["active"] else ""
            table.add_row(Text(f"{i}", style=f"bold {ACCENT}"), Text(p["name"], style="bold"),
                          Text(p["model"]), Text(f"{kind_label(p)} · {host(p['base_url'])}", style=DIM),
                          Text(mark, style="green"))
        console.print(Panel(table, title="API 配置", title_align="left", border_style=ACCENT))
        ans = plain_prompt(HTML("  <b>序号</b> 切换 · <b>a</b> 添加 · <b>e序号</b> 修改 · "
                                "<b>d序号</b> 删除 · 回车 返回 : ")).strip()
        if not ans:
            console.print()
            return active_profile()
        if ans.isdigit() and 1 <= int(ans) <= len(profiles):
            STORE["active"] = int(ans) - 1
            save_config()
            p = active_profile()
            console.print(Text(f"  ✓ 已切换到「{p['name']}」（{p['model']}），当前对话会接着聊\n", style="green"))
            return p
        if ans.lower() == "a":
            try:
                p = setup_wizard()
            except (KeyboardInterrupt, EOFError):
                console.print(Text("\n  已取消添加\n", style=DIM))
                continue
            profiles.append(p)
            STORE["active"] = len(profiles) - 1
            save_config()
            return p
        m = re.fullmatch(r"([eEdD])\s*(\d+)", ans)
        if m and 1 <= int(m.group(2)) <= len(profiles):
            i = int(m.group(2)) - 1
            if m.group(1).lower() == "e":
                try:
                    profiles[i] = setup_wizard(profiles[i])
                except (KeyboardInterrupt, EOFError):
                    console.print(Text("\n  已取消修改\n", style=DIM))
                    continue
                save_config()
                return active_profile()
            if len(profiles) == 1:
                console.print(Text("  只剩这一套了，不能删除（可以用 e1 修改它）", style="red"))
                continue
            if plain_prompt(HTML(f"  确定删除「{profiles[i]['name']}」吗？(y/n) : ")).strip().lower() in ("y", "yes", "是"):
                del profiles[i]
                if STORE["active"] >= i and STORE["active"] > 0:
                    STORE["active"] -= 1
                save_config()
                console.print(Text("  ✓ 已删除\n", style="green"))
            continue
        console.print(Text("  没看懂，请按提示输入", style="red"))


# ───────────────────────── 界面 ─────────────────────────

def logo_key(model):
    m = model.lower()
    if "claude" in m:
        return "claude"
    if "deepseek" in m:
        return "deepseek"
    if any(k in m for k in ("codex", "gpt", "openai", "o1", "o3", "o4")):
        return "codex"
    return "other"


# 主题色：边框、提示符、高亮文字都用它
THEMES = {"claude": "#da7758", "deepseek": "#7aa2ff", "codex": "#d4d4d8", "other": "#7aa2ff"}


def apply_theme(cfg):
    global ACCENT
    ACCENT = THEMES[logo_key(cfg["model"])]


def brand_name(model):
    return {"claude": "Claude", "deepseek": "DeepSeek", "codex": "Codex"}.get(
        logo_key(model), re.split(r"[-_/:.]", model)[0].capitalize())


def draw_logo(key):
    palette, rows = LOGOS[key]
    width = max(map(len, rows))
    rows = [r.ljust(width) for r in rows] + ([" " * width] if len(rows) % 2 else [])
    out = Text()
    for top, bottom in zip(rows[::2], rows[1::2]):
        for t, b in zip(top, bottom):
            if t == " " and b == " ":
                out.append(" ")
            elif b == " ":
                out.append("▀", style=palette[t])
            elif t == " ":
                out.append("▄", style=palette[b])
            else:
                out.append("▀", style=f"{palette[t]} on {palette[b]}")
        out.append("\n")
    out.rstrip()
    return out


def show_banner(cfg):
    apply_theme(cfg)
    console.clear()
    key = logo_key(cfg["model"])
    title = Text.assemble((brand_name(cfg["model"]) + "\n", f"bold {LOGOS[key][0][next(iter(LOGOS[key][0]))]}"),
                          (cfg["model"] + "\n", "bold"), (f"ds-chat {VERSION}", DIM))
    head = Table.grid(padding=(0, 4))
    logo = draw_logo(key)
    pad = max(0, (len(logo.plain.splitlines()) - 3) // 2)
    head.add_row(logo, Text("\n" * pad) + title)
    console.print(Padding(head, (1, 0, 1, 2)))
    info = Table.grid(padding=(0, 2))
    info.add_column(style=DIM)
    info.add_column()
    info.add_row("配置", Text.assemble((cfg["name"], f"bold {ACCENT}"), (f"   {kind_label(cfg)}", DIM)))
    info.add_row("模型", Text(cfg["model"], style=f"bold {ACCENT}"))
    info.add_row("线路", Text(host(cfg["base_url"])))
    info.add_row("模式", Text("Agent · 能上网、看图、读文档、运行命令、读写文件（改动前会问你）", style="yellow")
                 if STORE.get("agent", True) else Text("纯聊天", style=DIM))
    info.add_row("目录", Text(_short(os.getcwd())))
    tips = Text.assemble(
        ("提示  ", DIM),
        ("/api", ACCENT), (" 切换配置   ", DIM),
        ("/model", ACCENT), (" 切换模型   ", DIM),
        ("/his", ACCENT), (" 历史会话   ", DIM),
        ("/memory", ACCENT), (" 记忆   ", DIM),
        ("/clear", ACCENT), (" 新对话   ", DIM),
        ("/help", ACCENT), (" 帮助   ", DIM),
        ("/exit", ACCENT), (" 退出", DIM),
    )
    console.print(Panel(Group(info, Text(), tips), border_style=ACCENT,
                        title="[bold]欢迎[/]", title_align="left", padding=(1, 2)))
    console.print()


def show_help():
    table = Table(box=None, padding=(0, 2), show_header=False)
    for cmd, desc in [
        ("/api", "切换 / 添加 / 修改 / 删除 API 配置（比如 DeepSeek、Claude）"),
        ("/model", "切换模型"),
        ("/his", "历史会话：打开、删除一个、删除全部"),
        ("/memory", "AI 记住的关于你的事：查看、添加、删除"),
        ("/copy", "复制上一条回答"),
        ("/compact", "把前面的对话压缩成摘要（太长时也会自动压缩）"),
        ("/bar", "打开 / 关闭屏幕顶部菜单栏的 💰余额"),
        ("/agent", "打开 / 关闭 Agent 模式（AI 能否操作电脑）"),
        ("/line", "切换到备用线路（主线路连不上时用）"),
        ("/clear", "开始新对话（当前对话会留在历史里）"),
        ("/setup", "修改当前这套配置的线路、密钥、默认模型"),
        ("ai -c", "（在终端里运行）接着上次的对话聊"),
        ("ai --update", "（在终端里运行）更新到最新版"),
        ("/help", "显示这个帮助"),
        ("/exit", "退出（也可以按 Ctrl+D）"),
        ("拖图片进窗口", "把图片或截图拖进来，AI 就能看到"),
        ("Esc 然后 回车", "换行（写多行内容）"),
        ("Ctrl+C", "回答或操作时按：停止"),
    ]:
        table.add_row(Text(cmd, style=ACCENT), desc)
    console.print(Panel(table, title=f"帮助 · ds-chat {VERSION}", title_align="left", border_style=ACCENT), "\n")


# ───────────────────────── Agent 工具 ─────────────────────────

MAX_STEPS = 30          # 一次提问里最多连续调用多少次工具
MAX_OUTPUT = 12000      # 工具结果最多返回给 AI 多少字符
COMMAND_TIMEOUT = 120   # 命令最长运行秒数


def _tool(name, desc, props, required):
    return {"type": "function", "function": {
        "name": name, "description": desc,
        "parameters": {"type": "object", "properties": props, "required": required},
    }}


TOOLS = [
    _tool("run_command", "在用户的 Mac 上运行一条 shell 命令（zsh），返回输出。不能运行需要键盘交互的命令。",
          {"command": {"type": "string", "description": "要运行的命令"}}, ["command"]),
    _tool("read_file", "读取文件内容。支持文本文件，以及 PDF、Word（doc/docx）、Excel（xlsx）、PPT（pptx）、RTF 等文档。",
          {"path": {"type": "string", "description": "文件路径，可以用 ~ 表示主目录"}}, ["path"]),
    _tool("list_dir", "列出一个文件夹里的文件和子文件夹（含大小、修改时间）。",
          {"path": {"type": "string", "description": "文件夹路径，默认当前目录"}}, []),
    _tool("write_file", "创建或整个覆盖一个文件。",
          {"path": {"type": "string"}, "content": {"type": "string", "description": "完整的文件内容"}},
          ["path", "content"]),
    _tool("edit_file", "把文件里的一段文字替换成新文字（old_text 必须在文件里只出现一次）。",
          {"path": {"type": "string"}, "old_text": {"type": "string"}, "new_text": {"type": "string"}},
          ["path", "old_text", "new_text"]),
    _tool("web_search", "用搜索引擎搜索网页，返回标题、链接和摘要。查新闻、最新信息、报错原因、价格等都先用它。",
          {"query": {"type": "string", "description": "搜索关键词"}}, ["query"]),
    _tool("web_fetch", "打开一个网址，返回网页正文（去掉广告和代码）。搜索后想看详细内容时用。",
          {"url": {"type": "string", "description": "完整网址，http 或 https 开头"}}, ["url"]),
    _tool("view_image", "查看一张图片（png、jpg、heic、gif、webp 等，PDF 只看第一页）。",
          {"path": {"type": "string", "description": "图片路径"}}, ["path"]),
    _tool("screenshot", "截取用户整个屏幕并查看，用来了解用户正在看什么、某个 App 的状态。", {}, []),
    _tool("remember", "把关于用户的长期有用信息记下来（偏好、习惯、常用路径、项目情况等），以后的对话里都能看到。"
          "用户说“记住…”或透露了以后也有用的信息时使用。不要记密码、密钥这类敏感信息。",
          {"fact": {"type": "string", "description": "要记住的一句话"}}, ["fact"]),
    _tool("forget_memory", "删除一条记住的信息（编号见系统提示里的“你记住的事”）。",
          {"number": {"type": "string", "description": "要删除的编号，如 2"}}, ["number"]),
]

# 这些工具会改动电脑或看你的屏幕，执行前要用户同意
NEEDS_APPROVAL = {"run_command", "write_file", "edit_file", "screenshot"}


def system_prompt():
    home = os.path.expanduser("~")
    return (
        "你是运行在用户 Mac 终端里的 AI 助手，可以用工具查看文件、运行命令、写文件，帮用户把事情真正做完。\n"
        f"当前目录：{os.getcwd()}\n用户主目录：{home}\n今天：{time.strftime('%Y-%m-%d')}\n"
        "规则：\n"
        "- 用中文回答。用户是电脑新手，解释要通俗，少用术语。\n"
        "- 需要了解电脑上的情况时，自己用工具去查，不要让用户去运行命令。\n"
        "- 运行命令、写文件前，用户会看到并确认；被拒绝就换个办法或询问用户。\n"
        "- 删除、覆盖重要文件、sudo、修改系统设置这类有风险的操作，先说明风险再做。\n"
        "- 不要运行需要键盘交互或一直不结束的命令（如 vim、top、不带 -c 的 ping）。\n"
        f"- 每条命令都在 {os.getcwd()} 单独运行，上一条命令里的 cd 不会保留；需要换目录就在同一条命令里写 cd xxx && …。\n"
        "- 命令没有终端可以输入密码，sudo 之类要密码的命令会直接失败；需要管理员权限时，把命令给用户，让用户自己在终端运行。\n"
        "- 遇到需要最新信息、你不确定的事实、报错原因，先用 web_search 搜索，必要时用 web_fetch 打开网页看详情，"
        "回答时附上信息来源的网址。\n"
        "- 用户消息里如果附了图片，你能直接看到；需要看某张图片文件用 view_image，需要看用户屏幕用 screenshot。\n"
        "- 你可以通过命令操作这台 Mac：open -a \"App 名\" 打开 App，open 网址 用浏览器打开网页，"
        "osascript -e 'display notification \"内容\" with title \"标题\"' 弹通知，pbcopy 复制到剪贴板，say 朗读。\n"
        "- 用户说了以后也有用的个人信息或偏好（或者说“记住…”），用 remember 记下来；不要记密码和密钥。\n"
        "- 任务完成后，简短告诉用户做了什么、结果如何。"
        + memory_prompt()
    )


def _path(p):
    return os.path.abspath(os.path.expanduser(p or "."))


def _clip(text):
    if len(text) <= MAX_OUTPUT:
        return text
    half = MAX_OUTPUT // 2
    return text[:half] + f"\n\n…（中间省略 {len(text) - MAX_OUTPUT} 个字符）…\n\n" + text[-half:]


def _short(p):
    """把开头的主目录显示成 ~（只换开头，不换路径中间碰巧一样的部分）。"""
    home = os.path.expanduser("~").rstrip("/")
    if p == home or p.startswith(home + "/"):
        return "~" + p[len(home):]
    return p


def _kill_group(proc):
    try:
        os.killpg(proc.pid, signal.SIGKILL)
    except (ProcessLookupError, PermissionError):
        pass


def tool_run_command(command):
    # 单独的进程组：超时或按 Ctrl+C 时连同它启动的子进程一起结束。
    # 也没有控制终端，sudo 这类要输密码的命令会直接报错，而不是藏在转圈后面卡住。
    proc = subprocess.Popen(command, shell=True, executable="/bin/zsh", cwd=os.getcwd(),
                            stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                            text=True, errors="replace", start_new_session=True)
    try:
        stdout, stderr = proc.communicate(timeout=COMMAND_TIMEOUT)
    except subprocess.TimeoutExpired:
        _kill_group(proc)
        stdout, stderr = proc.communicate()
        return f"命令运行超过 {COMMAND_TIMEOUT} 秒，已停止。\n{stdout}{stderr}".strip(), False
    except KeyboardInterrupt:
        _kill_group(proc)
        proc.wait()
        raise
    out = (stdout or "") + (("\n[stderr]\n" + stderr) if stderr else "")
    return f"退出码 {proc.returncode}\n{out}".strip(), proc.returncode == 0


def tool_read_file(path):
    p = _path(path)
    if os.path.isdir(p):
        return f"{p} 是文件夹，请用 list_dir", False
    if os.path.splitext(p)[1].lower() in IMAGE_EXTS:
        return f"{p} 是图片，请用 view_image 查看", False
    doc = read_document(p)
    if doc is not None:
        return doc, True
    with open(p, "rb") as f:
        data = f.read(400_000)
    if b"\0" in data[:4000]:
        return f"{p} 是二进制文件（{os.path.getsize(p)} 字节），无法按文本读取", False
    return data.decode("utf-8", "replace"), True


def tool_list_dir(path=None):
    p = _path(path)
    rows = []
    for name in sorted(os.listdir(p), key=str.lower)[:500]:
        full = os.path.join(p, name)
        try:
            st = os.stat(full)
        except OSError:
            continue
        kind = "📁" if os.path.isdir(full) else "  "
        when = time.strftime("%Y-%m-%d %H:%M", time.localtime(st.st_mtime))
        rows.append(f"{kind} {name}  ({st.st_size} 字节, {when})")
    return f"{p}：\n" + ("\n".join(rows) or "（空文件夹）"), True


DOC_TEXTUTIL = {".doc", ".docx", ".rtf", ".rtfd", ".odt", ".html", ".htm", ".webarchive", ".wordml"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".heic", ".heif", ".bmp", ".tif", ".tiff"}


def read_document(p):
    """PDF、Word、Excel、PPT 等文档转成文字。返回 None 表示不是这些格式。"""
    ext = os.path.splitext(p)[1].lower()
    if ext == ".pdf":
        import logging
        logging.getLogger("pypdf").setLevel(logging.ERROR)  # 不然有些 PDF 会在聊天界面里打出一堆警告
        from pypdf import PdfReader
        reader = PdfReader(p)
        pages = []
        for i, page in enumerate(reader.pages[:60], 1):
            pages.append(f"--- 第 {i} 页 ---\n{(page.extract_text() or '').strip()}")
        text = "\n".join(pages)
        if not re.sub(r"--- 第 \d+ 页 ---|\s", "", text):
            return f"这个 PDF（共 {len(reader.pages)} 页）没有文字层，可能是扫描件或图片。可以用 view_image 看第一页。"
        more = f"\n…（共 {len(reader.pages)} 页，只读了前 60 页）" if len(reader.pages) > 60 else ""
        return text.translate(_LOOKALIKE) + more
    if ext in DOC_TEXTUTIL:
        r = subprocess.run(["/usr/bin/textutil", "-convert", "txt", "-stdout", p],
                           capture_output=True, text=True, errors="replace", timeout=60)
        if r.returncode != 0:
            raise RuntimeError(r.stderr.strip() or "无法读取这个文档")
        return r.stdout
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 2000:
                    rows.append("…（只显示前 2000 行）")
                    break
                if any(v is not None for v in row):
                    rows.append("\t".join("" if v is None else str(v) for v in row))
            parts.append(f"=== 工作表：{ws.title} ===\n" + "\n".join(rows))
        return "\n\n".join(parts)
    if ext == ".pptx":
        with zipfile.ZipFile(p) as z:
            slides = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                            key=lambda n: int(re.search(r"\d+", n.rsplit("/", 1)[1]).group()))
            out = []
            for i, n in enumerate(slides, 1):
                texts = re.findall(r"<a:t>([^<]*)</a:t>", z.read(n).decode("utf-8", "replace"))
                out.append(f"--- 第 {i} 页 ---\n" + html.unescape(" ".join(texts)))
        return "\n".join(out)
    if ext == ".xls":
        return "这是旧版 Excel（.xls），读不了。可以让用户用 Excel 或 Numbers 另存为 .xlsx。"
    return None


def load_image(path):
    """把图片缩小并转成 JPEG（省流量，HEIC 等格式模型也能看），返回 {"media_type", "data", "name"}。"""
    p = _path(path)
    if not os.path.isfile(p):
        raise FileNotFoundError(f"找不到图片：{p}")
    tmp_dir = os.path.join(APP_DIR, "tmp")
    os.makedirs(tmp_dir, exist_ok=True)
    out = os.path.join(tmp_dir, f"img-{os.getpid()}-{int(time.time() * 1000)}.jpg")
    try:
        r = subprocess.run(["/usr/bin/sips", "-s", "format", "jpeg", "-s", "formatOptions", "80",
                            "-Z", "1568", p, "--out", out], capture_output=True, timeout=60)
        if r.returncode != 0 or not os.path.exists(out):
            raise RuntimeError("这个文件不是图片，或者格式不支持")
        with open(out, "rb") as f:
            data = f.read()
    finally:
        if os.path.exists(out):
            os.remove(out)
    return {"media_type": "image/jpeg", "data": base64.b64encode(data).decode(), "name": os.path.basename(p)}


def tool_view_image(path):
    img = load_image(path)
    return f"已打开图片 {img['name']}（图片附在后面）", True, [img]


def tool_screenshot():
    tmp = os.path.join(APP_DIR, "tmp", f"screen-{os.getpid()}.png")
    os.makedirs(os.path.dirname(tmp), exist_ok=True)
    try:
        r = subprocess.run(["/usr/sbin/screencapture", "-x", "-m", tmp], capture_output=True, timeout=30)
        if r.returncode != 0 or not os.path.exists(tmp):
            return ("截屏失败。可能需要在 系统设置 → 隐私与安全性 → 屏幕录制 里允许「终端」。", False)
        img = load_image(tmp)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return ("已截取屏幕（图片附在后面）。如果画面里只有桌面背景、看不到窗口，说明还没有屏幕录制权限，"
            "请告诉用户在 系统设置 → 隐私与安全性 → 屏幕录制 里允许「终端」。", True, [img])


# ── 上网 ──

BROWSER_UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/130.0.0.0 Safari/537.36")


def _http_get(url, timeout=20, limit=3_000_000):
    req = urllib.request.Request(url, headers={"User-Agent": BROWSER_UA, "Accept-Encoding": "gzip",
                                               "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        data = resp.read(limit)
        ctype = resp.headers.get("Content-Type", "")
        final_url = resp.geturl()
    if data[:2] == b"\x1f\x8b":  # 有的网站不管怎样都返回压缩过的内容
        try:
            data = gzip.decompress(data)
        except (OSError, EOFError):
            pass
    m = re.search(r"charset=([\w-]+)", ctype) or re.search(rb'<meta[^>]+charset=["\']?([\w-]+)', data[:3000])
    charset = m.group(1) if m else "utf-8"
    if isinstance(charset, bytes):
        charset = charset.decode()
    try:
        text = data.decode(charset, "replace")
    except LookupError:
        text = data.decode("utf-8", "replace")
    return text, ctype, data, final_url


def _strip_tags(s):
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", s))).strip()


def _bing_url(href):
    """Bing 的结果链接有时是跳转链接，参数 u 里是 a1 + base64 编码的真实网址。"""
    href = html.unescape(href)
    if "bing.com/ck/a" in href:
        u = urllib.parse.parse_qs(urllib.parse.urlparse(href).query).get("u", [""])[0]
        if u.startswith("a1"):
            try:
                return base64.urlsafe_b64decode(u[2:] + "=" * (-len(u[2:]) % 4)).decode()
            except Exception:
                pass
    return href


def _search_bing(query):
    page = _http_get("https://cn.bing.com/search?" + urllib.parse.urlencode({"q": query}))[0]
    results = []
    for block in page.split('<li class="b_algo"')[1:]:
        a = re.search(r'<h2[^>]*>\s*<a[^>]+href="([^"]+)"[^>]*>(.*?)</a>', block, re.S)
        if not a:
            continue
        snip = re.search(r'<p[^>]*>(.*?)</p>', block, re.S)
        results.append((_strip_tags(a.group(2)), _bing_url(a.group(1)), _strip_tags(snip.group(1)) if snip else ""))
    return results


def _search_360(query):
    page = _http_get("https://www.so.com/s?" + urllib.parse.urlencode({"q": query}))[0]
    results = []
    heads = list(re.finditer(r'<h3 class="res-title[^"]*"[^>]*>\s*<a([^>]*)>(.*?)</a>', page, re.S))
    for i, m in enumerate(heads):
        real = re.search(r'data-mdurl="([^"]+)"', m.group(1)) or re.search(r'href="([^"]+)"', m.group(1))
        if not real:
            continue
        end = heads[i + 1].start() if i + 1 < len(heads) else m.end() + 3000
        snippet = _strip_tags(re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", page[m.end():end]))
        results.append((_strip_tags(m.group(2)), html.unescape(real.group(1)), snippet[:200]))
    return results


def tool_web_search(query):
    errors = []
    for name, engine in (("Bing", _search_bing), ("360", _search_360)):
        try:
            results = engine(query)[:8]
        except Exception as e:
            errors.append(f"{name}：{e}")
            continue
        if results:
            lines = [f"{i}. {t}\n   {u}" + (f"\n   {s[:200]}" if s else "")
                     for i, (t, u, s) in enumerate(results, 1)]
            return f"搜索「{query}」（{name}）：\n" + "\n".join(lines), True
        errors.append(f"{name}：没有结果")
    return "搜索失败：" + "；".join(errors), False


def tool_web_fetch(url):
    if not re.match(r"https?://", url):
        return "网址要以 http:// 或 https:// 开头", False
    text, ctype, data, final_url = _http_get(url)
    for _ in range(3):  # 跟随搜索结果里的 JS / meta 跳转页
        if len(text) > 5000:
            break
        m = (re.search(r'(?:window\.)?location(?:\.href)?\s*(?:=|\.replace\()\s*["\']([^"\']+)["\']', text)
             or re.search(r'<meta[^>]+http-equiv=["\']?refresh["\']?[^>]+url=([^"\'>]+)', text, re.I))
        if not m:
            break
        nxt = urllib.parse.urljoin(final_url, html.unescape(m.group(1)).strip("'\""))
        if nxt == final_url or not nxt.startswith("http"):
            break
        text, ctype, data, final_url = _http_get(nxt)
    if "pdf" in ctype or data[:5] == b"%PDF-":
        tmp = os.path.join(APP_DIR, "tmp", f"web-{os.getpid()}.pdf")
        os.makedirs(os.path.dirname(tmp), exist_ok=True)
        with open(tmp, "wb") as f:
            f.write(data)
        try:
            return read_document(tmp), True
        finally:
            os.remove(tmp)
    if "html" not in ctype and "<html" not in text[:2000].lower():
        return text, True
    title = re.search(r"<title[^>]*>(.*?)</title>", text, re.S | re.I)
    body = re.sub(r"(?is)<(script|style|noscript|svg|nav|footer|header|aside|form|iframe)\b.*?</\1>", " ", text)
    body = re.sub(r"(?is)<!--.*?-->", " ", body)
    body = re.sub(r"(?i)<(br|/p|/div|/li|/h[1-6]|/tr|/section|/article)[^>]*>", "\n", body)
    body = html.unescape(re.sub(r"<[^>]+>", " ", body))
    lines = [re.sub(r"[ \t\u3000\xa0]+", " ", l).strip() for l in body.split("\n")]
    body = "\n".join(l for l in lines if l)
    head = f"网页：{_strip_tags(title.group(1)) if title else ''}\n网址：{final_url}\n\n"
    return head + (body or "（这个网页没有可读的文字，可能需要浏览器运行脚本才能显示）"), True


# ── 长期记忆 ──

MEMORY_FILE = os.path.join(CONFIG_DIR, "memory.json")


def load_memory():
    try:
        with open(MEMORY_FILE, encoding="utf-8") as f:
            items = json.load(f)
        return [m for m in items if isinstance(m, dict) and m.get("text")]
    except (FileNotFoundError, json.JSONDecodeError, TypeError):
        return []


def save_memory(items):
    os.makedirs(CONFIG_DIR, exist_ok=True)
    with open(MEMORY_FILE, "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False, indent=2)
    os.chmod(MEMORY_FILE, 0o600)


def memory_prompt():
    items = load_memory()
    if not items:
        return ""
    return "\n\n你记住的关于用户的事（来自以前的对话）：\n" + "\n".join(
        f"{i}. {m['text']}" for i, m in enumerate(items, 1))


def tool_remember(fact):
    items = load_memory()
    fact = fact.strip()[:300]
    if any(m["text"] == fact for m in items):
        return "已经记住过了", True
    items.append({"text": fact, "time": time.strftime("%Y-%m-%d")})
    save_memory(items)
    return f"已记住（第 {len(items)} 条）", True


def tool_forget_memory(number):
    items = load_memory()
    if not number.strip().isdigit() or not 1 <= int(number) <= len(items):
        return f"没有第 {number} 条", False
    gone = items.pop(int(number) - 1)
    save_memory(items)
    return f"已删除：{gone['text']}", True


def tool_write_file(path, content):
    p = _path(path)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "w", encoding="utf-8") as f:
        f.write(content)
    return f"已写入 {p}（{len(content.splitlines())} 行）", True


def tool_edit_file(path, old_text, new_text):
    p = _path(path)
    with open(p, encoding="utf-8") as f:
        text = f.read()
    n = text.count(old_text)
    if n != 1:
        return f"old_text 在文件里出现了 {n} 次，必须正好 1 次，请提供更长、唯一的片段", False
    with open(p, "w", encoding="utf-8") as f:
        f.write(text.replace(old_text, new_text, 1))
    return f"已修改 {p}", True


TOOL_FUNCS = {
    "run_command": tool_run_command, "read_file": tool_read_file, "list_dir": tool_list_dir,
    "write_file": tool_write_file, "edit_file": tool_edit_file,
    "web_search": tool_web_search, "web_fetch": tool_web_fetch, "view_image": tool_view_image,
    "screenshot": tool_screenshot, "remember": tool_remember, "forget_memory": tool_forget_memory,
}


def preview_tool(name, args):
    """执行前在屏幕上展示 AI 要做什么。"""
    if name == "run_command":
        console.print(Panel(Text("$ " + args.get("command", ""), style="bold"),
                            title="⚡ 运行命令", title_align="left", border_style="yellow"))
    elif name == "write_file":
        content = args.get("content", "")
        lines = content.splitlines()
        body = "\n".join(lines[:15]) + (f"\n…（共 {len(lines)} 行）" if len(lines) > 15 else "")
        exists = "覆盖" if os.path.exists(_path(args.get("path"))) else "新建"
        console.print(Panel(Text(body, style=DIM),
                            title=Text(f"✎ {exists}文件 {_short(_path(args.get('path')))}"),
                            title_align="left", border_style="yellow"))
    elif name == "edit_file":
        diff = Text()
        for l in args.get("old_text", "").splitlines()[:12]:
            diff.append("- " + l + "\n", style="red")
        for l in args.get("new_text", "").splitlines()[:12]:
            diff.append("+ " + l + "\n", style="green")
        console.print(Panel(diff, title=Text(f"✎ 修改文件 {_short(_path(args.get('path')))}"),
                            title_align="left", border_style="yellow"))
    elif name == "read_file":
        console.print(Text(f"  📄 读取 {_short(_path(args.get('path')))}", style=DIM))
    elif name == "list_dir":
        console.print(Text(f"  📁 查看 {_short(_path(args.get('path')))}", style=DIM))
    elif name == "web_search":
        console.print(Text(f"  🌐 搜索：{args.get('query', '')}", style=DIM))
    elif name == "web_fetch":
        console.print(Text(f"  🌐 打开 {args.get('url', '')}", style=DIM))
    elif name == "view_image":
        console.print(Text(f"  🖼  查看图片 {_short(_path(args.get('path')))}", style=DIM))
    elif name == "screenshot":
        console.print(Panel(Text("截取整个屏幕给 AI 看", style="bold"),
                            title="📸 截屏", title_align="left", border_style="yellow"))
    elif name == "remember":
        console.print(Text(f"  🧠 记住：{args.get('fact', '')}", style="magenta"))
    elif name == "forget_memory":
        console.print(Text(f"  🧠 忘掉第 {args.get('number', '')} 条", style="magenta"))


def approve(name, state):
    """返回 True 允许 / False 拒绝；Ctrl+C 往外抛，结束这一轮。"""
    if name not in NEEDS_APPROVAL or state.get("auto_approve"):
        return True
    while True:
        with VIEW.paused():
            ans = plain_prompt(HTML(
                "  <b>允许吗？</b> <ansigreen>y</ansigreen> 允许 · "
                "<ansicyan>a</ansicyan> 本次打开期间都允许 · <ansired>n</ansired> 拒绝 : ")).strip().lower()
        if ans in ("y", "yes", "是"):
            return True
        if ans in ("a", "all"):
            state["auto_approve"] = True
            return True
        if ans in ("n", "no", "否"):
            return False
        console.print(Text("  请输入 y（允许）、a（都允许）或 n（拒绝）", style="red"))


# 模型偶尔会用别的工具里常见的参数名（比如 file_path），换成我们的名字
PARAM_ALIASES = {"file_path": "path", "filepath": "path", "file": "path", "filename": "path",
                 "dir": "path", "directory": "path", "folder": "path", "cmd": "command",
                 "old_string": "old_text", "new_string": "new_text", "old": "old_text", "new": "new_text",
                 "text": "content", "data": "content"}


def _warn(msg):
    console.print(Text(f"  ⚠ {msg}", style="red"))
    return msg


class Denied(Exception):
    """用户拒绝了操作：整轮停下，等用户说怎么做。"""


def execute_tool(call, state):
    """执行一次工具调用，返回 (给 AI 的文字结果, 图片列表)。"""
    text, images = _execute_tool(call, state)
    return _clip(text), images


def _execute_tool(call, state):
    name = call["function"]["name"]
    try:
        args = json.loads(call["function"].get("arguments") or "{}")
    except json.JSONDecodeError:
        return (_warn("AI 给的操作参数格式不对，已让它重试"), [])
    if not isinstance(args, dict):
        return (_warn("AI 给的操作参数格式不对，已让它重试"), [])
    func = TOOL_FUNCS.get(name)
    if func is None:
        return (_warn(f"AI 想用一个不存在的工具 {name}，已让它重试"), [])
    schema = next(t["function"]["parameters"] for t in TOOLS if t["function"]["name"] == name)
    for k in list(args):
        alias = PARAM_ALIASES.get(k)
        if k not in schema["properties"] and alias in schema["properties"] and alias not in args:
            args[alias] = args.pop(k)
    bad = [k for k in schema["required"] if not isinstance(args.get(k), str)]
    if bad:
        return (_warn(f"AI 调用 {name} 时缺少参数 {', '.join(bad)}，已让它重试"), [])
    args = {k: v for k, v in args.items() if k in schema["properties"]}
    preview_tool(name, args)
    if not approve(name, state):
        console.print(Text("  ✗ 已拒绝，AI 已停下。你可以告诉它换个做法。\n", style="red"))
        raise Denied
    busy = {"run_command": "运行中…", "web_search": "搜索中…", "web_fetch": "打开网页中…",
            "screenshot": "截屏中…", "read_file": "读取中…"}.get(name)
    try:
        if busy:
            with VIEW.busy(busy):
                out = func(**args)
        else:
            out = func(**args)
    except TypeError as e:
        return f"参数不对：{e}", []
    except Exception as e:
        out = (f"出错：{e}", False)
    result, ok = out[0], out[1]
    images = out[2] if len(out) > 2 else []
    if name in NEEDS_APPROVAL:
        lines = result.splitlines()
        shown = "\n".join(lines[-12:]) if len(lines) > 12 else result
        if len(lines) > 12:
            shown = f"…（只显示最后 12 行，共 {len(lines)} 行）\n" + shown
        console.print(Padding(Text(shown, style="green" if ok else "red"), (0, 0, 1, 4)))
    elif not ok:
        console.print(Padding(Text(result[:500], style="red"), (0, 0, 0, 4)))
    return result, images


# ───────────────────────── 对话 ─────────────────────────

def stream_chat(base, key, model, messages, tools=None):
    """逐块产出 ("reasoning"|"content"|"tool"|"usage", 数据)。"""
    payload = {
        "model": model,
        "messages": messages,
        "stream": True,
        "stream_options": {"include_usage": True},
    }
    if tools:
        payload["tools"] = tools
    req = urllib.request.Request(
        f"{base}/chat/completions", data=json.dumps(payload).encode(), method="POST",
        headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json",
                 "User-Agent": USER_AGENT},
    )
    got_data, other = False, []
    with urllib.request.urlopen(req, timeout=300) as resp:
        for raw in resp:
            line = raw.decode("utf-8", "replace").strip()
            if not line.startswith("data:"):
                if line and sum(map(len, other)) < 4000:
                    other.append(line)
                continue
            got_data = True
            data = line[5:].strip()
            if data == "[DONE]":
                break
            try:
                chunk = json.loads(data)
            except json.JSONDecodeError:
                continue
            if chunk.get("error"):
                raise RuntimeError(f"接口报错：{_error_text(chunk['error'])}")
            if chunk.get("usage"):
                yield "usage", chunk["usage"]
            for choice in chunk.get("choices") or []:
                delta = choice.get("delta") or {}
                if delta.get("reasoning_content"):
                    yield "reasoning", delta["reasoning_content"]
                if delta.get("content"):
                    yield "content", delta["content"]
                for tc in delta.get("tool_calls") or []:
                    yield "tool", tc
    if not got_data:
        # 中转站有时用 200 状态码返回一段错误 JSON，而不是流式数据
        body = "\n".join(other)
        try:
            body = _error_text(json.loads(body).get("error") or body)
        except (json.JSONDecodeError, AttributeError):
            pass
        raise RuntimeError(f"接口没有返回内容：{body[:300] or '（空）'}")


def _error_text(err):
    if isinstance(err, dict):
        return str(err.get("message") or err)
    return str(err)


def fetch_balance(cfg):
    """查余额（美元）。支持两种中转站接口；查不到返回 None。"""
    if is_official(cfg):
        return None  # Claude 官方接口没有查余额的地方
    v1 = cfg["base_url"].rstrip("/")
    if not v1.endswith("/v1"):
        v1 += "/v1"

    def get(path):
        req = urllib.request.Request(v1 + path, headers={"Authorization": f"Bearer {cfg['api_key']}",
                                                          "User-Agent": USER_AGENT})
        with urllib.request.urlopen(req, timeout=15) as resp:
            return json.load(resp)

    try:
        return float(get("/usage")["balance"])
    except Exception:
        pass
    try:  # new-api / one-api：总额度 - 已用（已用的单位是美分）
        limit = float(get("/dashboard/billing/subscription")["hard_limit_usd"])
        used = float(get("/dashboard/billing/usage")["total_usage"]) / 100
        return None if limit >= 1e7 else limit - used  # 额度“无限”时算不出余额
    except Exception:
        return None


def refresh_balance(cfg, stats, before=None):
    """后台查余额，结果写进 stats["balance"]（查不到为 "n/a"）。
    传入 before（这次提问前的余额）时，顺便算出这次花了多少，写进 stats["last_cost"]。"""
    def run():
        bal = fetch_balance(cfg)
        if bal is None:
            stats["balance"] = "n/a"
        else:
            stats["balance"] = bal
            if isinstance(before, float):
                stats["last_cost"] = max(before - bal, 0.0)
        MENUBAR.update(stats["balance"])
    threading.Thread(target=run, daemon=True).start()


class MenuBar:
    """屏幕顶部菜单栏显示余额（橘色钱袋图标）（ai 打开时显示，退出后消失）。
    显示由一个很小的原生程序 dsbar 负责：它每 2 秒读一次 menubar.txt，文件超过 20 秒没更新就自己退出。
    开着几个 ai 窗口也只有一个 dsbar。"""
    BIN = os.path.join(APP_DIR, "dsbar")
    FILE = os.path.join(APP_DIR, "menubar.txt")
    PID = os.path.join(APP_DIR, "dsbar.pid")

    def __init__(self):
        self.text, self.on = "", False

    def start(self):
        if not STORE.get("menubar", True) or not os.access(self.BIN, os.X_OK) or self.on:
            return
        self.on = True
        self._write()
        if not self._running():
            proc = subprocess.Popen([self.BIN, self.FILE], stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                                    stderr=subprocess.DEVNULL, start_new_session=True)
            with open(self.PID, "w") as f:
                f.write(str(proc.pid))
        threading.Thread(target=self._heartbeat, daemon=True).start()

    def _running(self):
        try:
            with open(self.PID) as f:
                pid = int(f.read().strip())
            os.kill(pid, 0)
            comm = subprocess.run(["ps", "-p", str(pid), "-o", "comm="], capture_output=True, text=True).stdout
            return comm.strip().endswith("dsbar")
        except (OSError, ValueError):
            return False

    def _heartbeat(self):
        while self.on:
            self._write()
            time.sleep(5)

    def _write(self):
        try:
            with open(self.FILE, "w", encoding="utf-8") as f:
                f.write(self.text if self.on else "")
        except OSError:
            pass

    def update(self, balance):
        self.text = f"{balance:.2f}" if isinstance(balance, float) else ""  # 图标由 dsbar 自己画
        if self.on:
            self._write()

    def stop(self):
        self.on = False
        self._write()  # 写空：菜单栏马上隐藏，dsbar 20 秒后自己退出（别的 ai 窗口还开着的话会继续写）


MENUBAR = MenuBar()


def md(text):
    """Markdown 默认把单个换行当空格，聊天里会把一行一行的内容挤成一段；改成换行就是换行（代码块不动）。"""
    out, fence = [], False
    for line in text.split("\n"):
        if line.lstrip().startswith(("```", "~~~")):
            fence = not fence
        elif not fence and line.strip() and not line.endswith("  "):
            line += "  "
        out.append(line)
    return Markdown("\n".join(out), code_theme="monokai")


def split_think(content):
    """有些模型把思考过程放在 <think>…</think> 里，拆出来。"""
    m = re.match(r"\s*<think>(.*?)(</think>|$)(.*)", content, re.S)
    if not m:
        return "", content
    return m.group(1), m.group(3) if m.group(2) else ""


def render(reasoning, answer, started, thinking_done_at, preparing_tool=False):
    parts = []
    if preparing_tool and not answer:
        parts.append(Spinner("dots", text=Text(" 正在准备操作…", style=f"italic {ACCENT}")))
    elif reasoning and not answer:
        tail = reasoning.strip().splitlines()[-1][-80:] if reasoning.strip() else ""
        parts.append(Spinner("dots", text=Text.assemble(
            (f" 思考中 {time.time() - started:.0f}s  ", f"italic {ACCENT}"), (tail, f"italic {DIM}"))))
    elif not answer:
        parts.append(Spinner("dots", text=Text(" 正在连接…", style=f"italic {ACCENT}")))
    if reasoning and (answer or preparing_tool):
        secs = (thinking_done_at or time.time()) - started
        parts.append(Text(f"✻ 已思考 {secs:.1f} 秒", style=f"italic {DIM}"))
    if answer:
        parts.append(md(answer))
    return Padding(Group(*parts), (0, 0, 0, 2))


class _Tail:
    """只显示最后 height 行。内容比窗口还高时，Live 刷新会把内容重复打印到屏幕上。"""

    def __init__(self, renderable, height):
        self.renderable, self.height = renderable, height

    def __rich_console__(self, console, options):
        lines = console.render_lines(self.renderable, options, pad=False)
        for line in lines[-self.height:]:
            yield from line
            yield Segment.line()


def _openai_image(img):
    return {"type": "image_url", "image_url": {"url": f"data:{img['media_type']};base64,{img['data']}"}}


_PATH_RE = re.compile(r"""'([^']+)'|"([^"]+)"|((?:~|/)(?:\\.|[^\s\\])+)""")


def attach_images(text):
    """找出输入里的图片路径（把图片拖进终端窗口就会变成路径），读成图片。返回 (图片列表, 出错信息列表)。"""
    images, errors, seen = [], [], set()
    for m in _PATH_RE.finditer(text):
        raw = m.group(1) or m.group(2) or re.sub(r"\\(.)", r"\1", m.group(3))
        path = os.path.expanduser(raw)
        if os.path.splitext(path)[1].lower() not in IMAGE_EXTS or path in seen or not os.path.isfile(path):
            continue
        seen.add(path)
        try:
            images.append(load_image(path))
        except Exception as e:
            errors.append(f"{os.path.basename(path)}：{e}")
    return images, errors


def estimate_tokens(history):
    """粗略估算对话有多少 token：中文大约 1 字 1 token，英文约 4 字母 1 token，这里取 2 个字符 1 token；图片按 1500 算。"""
    n_img = sum(len(m.get("images") or []) for m in history)
    slim = [{k: v for k, v in m.items() if k != "images"} for m in history]
    return len(json.dumps(slim, ensure_ascii=False)) // 2 + n_img * 1500


def context_limit(cfg):
    return 150_000 if "claude" in cfg["model"].lower() else 60_000


def _complete(cfg, system, prompt):
    """不带工具、不显示过程地问一次，返回回答文字（压缩对话用）。"""
    msgs = [{"role": "user", "content": prompt}]
    if cfg.get("type") == "anthropic":
        events = stream_anthropic(cfg, system, msgs)
    else:
        events = stream_chat(cfg["base_url"], cfg["api_key"], cfg["model"],
                             [{"role": "system", "content": system}] + msgs)
    return split_think("".join(d for k, d in events if k == "content"))[1].strip()


def compact(cfg, history):
    """把较早的对话压缩成一段摘要，只保留最近的部分。成功返回 True。"""
    users = [i for i, m in enumerate(history) if m["role"] == "user"]
    if len(users) < 2:
        return False
    split = users[-2] if estimate_tokens(history[users[-2]:]) < context_limit(cfg) * 0.3 else users[-1]
    lines = []
    for m in history[:split]:
        if m["role"] == "user":
            lines.append(f"用户：{m['content'][:3000]}")
        elif m["role"] == "assistant":
            if m.get("content"):
                lines.append(f"AI：{m['content'][:3000]}")
            for c in m.get("tool_calls") or []:
                lines.append(f"AI 调用工具 {c['function']['name']}：{c['function'].get('arguments', '')[:300]}")
        elif m["role"] == "tool":
            lines.append(f"工具结果：{m['content'][:400]}")
    transcript = "\n".join(lines)
    if len(transcript) > context_limit(cfg):
        half = context_limit(cfg) // 2
        transcript = transcript[:half] + "\n…（中间省略）…\n" + transcript[-half:]
    summary = _complete(cfg, "你负责把对话记录压缩成摘要。", (
        "下面是一段用户和 AI 助手的对话记录。请压缩成一份摘要，供 AI 接着对话时参考。\n"
        "要保留：用户的目标和要求、已经做过的操作和结果、重要的文件路径/网址/数字/结论、还没完成的事。\n"
        "用中文，分条写，尽量简短。\n\n" + transcript))
    if not summary:
        return False
    history[:split] = [{"role": "user", "content": "【之前对话的摘要】\n" + summary},
                       {"role": "assistant", "content": "好的，我已了解之前的对话内容，我们继续。"}]
    return True


def is_context_error(e):
    """接口报“内容太长、超过上下文”之类的错误。"""
    if error_code(e) not in (400, 413):
        return False
    if hasattr(e, "read") and not hasattr(e, "ds_body"):
        try:
            e.ds_body = e.read().decode("utf-8", "replace")
        except Exception:
            e.ds_body = ""
    text = (getattr(e, "ds_body", "") + str(getattr(e, "message", "")) + str(e)).lower()
    return any(w in text for w in ("context", "too long", "too many tokens", "maximum", "exceed", "长度", "上下文"))


def to_openai_messages(history):
    """发给 OpenAI 格式接口的消息：去掉 Claude 专用字段；之前几轮的思考内容也不再发回（省 token，有的模型不接受）。
    图片：用户消息里的图片放进 content 列表；工具返回的图片（工具消息只能是文字）在这批工具结果后面补一条用户消息。"""
    last_user = max(i for i, m in enumerate(history) if m["role"] == "user")
    out, pending = [], []
    for i, m in enumerate(history):
        if pending and m["role"] != "tool":
            out.append({"role": "user", "content": [{"type": "text", "text": "（上面工具返回的图片）"}]
                        + [_openai_image(x) for x in pending]})
            pending = []
        msg = {k: v for k, v in m.items() if k not in ("anthropic_content", "images")
               and (k != "reasoning_content" or i >= last_user)}
        if m.get("images"):
            if m["role"] == "user":
                msg["content"] = [{"type": "text", "text": m["content"]}] + [_openai_image(x) for x in m["images"]]
            else:
                pending += m["images"]
        out.append(msg)
    if pending:
        out.append({"role": "user", "content": [{"type": "text", "text": "（上面工具返回的图片）"}]
                    + [_openai_image(x) for x in pending]})
    return out


class TypeAhead:
    """AI 工作时在后台读键盘：打的字显示在底部输入框里，回车就排进队列，AI 做完当前这一步就会看到。"""

    def __init__(self):
        self.buf, self.queue = "", []
        self.lock = threading.Lock()
        self._thread, self._run, self._fd, self._old = None, False, None, None

    def start(self):
        if self._run or not sys.stdin.isatty():
            return
        import termios
        import tty
        self._fd = sys.stdin.fileno()
        self._old = termios.tcgetattr(self._fd)
        tty.setcbreak(self._fd)  # 一个字一个字地读、不回显；Ctrl+C 照样能停
        self._run = True
        self._thread = threading.Thread(target=self._loop, daemon=True)
        self._thread.start()

    def stop(self):
        if not self._run:
            return
        import termios
        self._run = False
        self._thread.join(1)
        termios.tcsetattr(self._fd, termios.TCSADRAIN, self._old)

    @property
    def active(self):
        return self._run

    def _loop(self):
        import codecs
        decode = codecs.getincrementaldecoder("utf-8")("ignore").decode
        esc = ""
        while self._run:
            if not select.select([self._fd], [], [], 0.1)[0]:
                continue
            for ch in decode(os.read(self._fd, 1024)):
                if esc:  # 跳过方向键之类的按键（ESC [ … 字母）
                    esc += ch
                    if (len(esc) > 1 and ch != "[" and (ch.isalpha() or ch == "~")) or len(esc) > 8:
                        esc = ""
                    continue
                with self.lock:
                    if ch == "\x1b":
                        esc = ch
                    elif ch in "\r\n":
                        if self.buf.strip():
                            self.queue.append(self.buf.strip())
                        self.buf = ""
                    elif ch in "\x7f\x08":
                        self.buf = self.buf[:-1]
                    elif ch == "\x15":  # Ctrl+U 清空
                        self.buf = ""
                    elif ch >= " ":
                        self.buf += ch

    def take(self):
        with self.lock:
            items, self.queue = self.queue, []
        return items

    def leftover(self):
        """没发出去的内容（排队的 + 正在打的），回到下一次的输入框里。"""
        with self.lock:
            text = "\n".join(self.queue + ([self.buf] if self.buf.strip() else []))
            self.queue, self.buf = [], ""
        return text


class TurnView:
    """AI 回答、执行操作期间屏幕底部的显示：正在生成的内容、转圈状态、补充输入框。"""

    def __init__(self):
        self.live, self.stream, self.status = None, None, None
        self.spinner = Spinner("dots")
        self.typeahead = TypeAhead()

    def __rich_console__(self, console, options):
        height = options.max_height or console.height
        box = self._box() if self.typeahead.active else None
        reserve = (3 if box else 0) + (1 if self.status else 0) + 1
        if self.stream is not None:
            yield _Tail(self.stream, max(height - reserve, 3))
        if self.status:
            self.spinner.text = Text(" " + self.status, style=f"italic {ACCENT}")
            yield Padding(self.spinner, (0, 0, 0, 2))
        if box:
            yield box

    def _box(self):
        with self.typeahead.lock:
            buf, queued = self.typeahead.buf, len(self.typeahead.queue)
        if buf:
            body = Text.assemble(("❯ ", f"bold {ACCENT}"), buf, ("▌", ACCENT))
        else:
            body = Text.assemble(("❯ ", f"bold {ACCENT}"), ("▌", ACCENT),
                                 ("  可以继续输入补充内容，回车发送，AI 做完这一步就会看到", DIM))
        title = f"已排队 {queued} 条补充" if queued else None
        return Panel(body, border_style=DIM, title=title, title_align="left", height=3)

    def start(self):
        self.typeahead.start()
        self.live = Live(self, console=console, refresh_per_second=12, transient=True, vertical_overflow="crop")
        self.live.start()

    def stop(self):
        if self.live:
            self.live.stop()
            self.live = None
        self.typeahead.stop()
        self.stream = self.status = None

    @contextmanager
    def paused(self):
        """要用普通输入框问用户（比如“允许吗？”）时，先把底部显示和后台读键盘停掉。"""
        if not self.live:
            yield
            return
        self.live.stop()
        self.typeahead.stop()
        try:
            yield
        finally:
            self.typeahead.start()
            self.live.start()

    @contextmanager
    def busy(self, text):
        """显示“运行中…”之类的转圈。底部显示开着时放进里面，否则用普通的转圈。"""
        if not self.live:
            with console.status(text, spinner="dots"):
                yield
            return
        old, self.status = self.status, text
        try:
            yield
        finally:
            self.status = old


VIEW = TurnView()


def stream_turn(cfg, history, stats):
    """请求一次模型并实时显示。返回 (回答, 思考内容, 工具调用列表, 额外信息)。Ctrl+C 往外抛。"""
    started = time.time()
    raw_content, reasoning_field = "", ""
    calls = {}  # index -> {"id", "type", "function": {"name", "arguments"}}
    extra = {}  # Claude 才有：工具调用、原始内容块、结束原因
    thinking_done_at = None
    tools = TOOLS if STORE.get("agent", True) else None
    interrupted = False
    if cfg.get("type") == "anthropic":
        events = stream_anthropic(cfg, system_prompt(), history, tools)
    else:
        messages = [{"role": "system", "content": system_prompt()}] + to_openai_messages(history)
        events = stream_chat(cfg["base_url"], cfg["api_key"], cfg["model"], messages, tools)

    # 流式输出显示在底部区域（只显示末尾几行），结束后清掉、再把完整内容打印一次
    VIEW.stream = render("", "", started, None)
    try:
        for kind, data in events:
            if kind == "final":
                extra = data
                continue
            if kind == "tool_pending":
                calls.setdefault(-1, None)  # 只是让界面显示“正在准备操作”
            elif kind == "reasoning":
                reasoning_field += data
            elif kind == "content":
                raw_content += data
            elif kind == "usage":
                stats["tokens"] += data.get("total_tokens", 0)
            elif kind == "tool":
                c = calls.setdefault(data.get("index", 0),
                                     {"id": "", "type": "function", "function": {"name": "", "arguments": ""}})
                c["id"] = data.get("id") or c["id"]
                fn = data.get("function") or {}
                c["function"]["name"] += fn.get("name") or ""
                c["function"]["arguments"] += fn.get("arguments") or ""
            think_tag, answer = split_think(raw_content)
            reasoning = reasoning_field + think_tag
            if (answer or calls) and thinking_done_at is None:
                thinking_done_at = time.time()
            VIEW.stream = render(reasoning, answer, started, thinking_done_at, bool(calls))
    except KeyboardInterrupt:
        interrupted = True
    finally:
        VIEW.stream = None
    think_tag, answer = split_think(raw_content)
    final = []
    if reasoning_field + think_tag:
        secs = (thinking_done_at or time.time()) - started
        final.append(Text(f"✻ 已思考 {secs:.1f} 秒", style=f"italic {DIM}"))
    if answer:
        final.append(md(answer))
    if interrupted:
        final.append(Text("（已中断）", style=DIM))
    if final:
        console.print(Padding(Group(*final), (0, 0, 0, 2)))
    if interrupted:
        raise KeyboardInterrupt
    if extra:
        return answer.strip(), reasoning_field + think_tag, extra["calls"], extra
    return answer.strip(), reasoning_field + think_tag, [calls[i] for i in sorted(calls) if i >= 0], {}


def take_supplements(history):
    """把用户在 AI 工作时补充输入的内容加进对话。返回加了几条。"""
    items = VIEW.typeahead.take()
    for text in items:
        console.print(Panel(Text(text), border_style=DIM, title="你（补充）", title_align="left"))
        history.append({"role": "user", "content": text})
    return len(items)


def ask(cfg, history, stats, state):
    """回答一次提问（中间可能调用很多次工具）。返回用户补充了但还没发出去的内容。"""
    started = time.time()
    tokens_before = stats["tokens"]
    VIEW.start()
    try:
        retried = False
        for _ in range(MAX_STEPS):
            take_supplements(history)
            if estimate_tokens(history) > context_limit(cfg):
                with VIEW.busy("对话太长了，正在把前面的内容压缩成摘要…"):
                    done = compact(cfg, history)
                if done:
                    console.print(Text("  🗜  对话太长，已把前面的内容压缩成摘要（输入 /his 打开旧会话仍能看到原文）", style=DIM))
            try:
                answer, reasoning, calls, extra = stream_turn(cfg, history, stats)
            except Exception as e:
                if retried or not is_context_error(e):
                    raise
                with VIEW.busy("内容超过模型上限，正在压缩前面的对话后重试…"):
                    done = compact(cfg, history)
                if not done:
                    raise
                retried = True
                console.print(Text("  🗜  内容超过模型上限，已压缩前面的对话并重试", style=DIM))
                answer, reasoning, calls, extra = stream_turn(cfg, history, stats)
            stop = extra.get("stop_reason")
            if stop == "refusal":
                cat = extra.get("refusal_category")
                console.print(Text(f"  Claude 拒绝回答这个请求{f'（类别：{cat}）' if cat else ''}，可以换个说法再试。",
                                   style="yellow"))
            if stop == "max_tokens" and calls:
                console.print(Text("  回答太长被截断了，这次的操作没有执行。可以说“继续”。", style="yellow"))
                calls = []
            if not calls:
                if answer or extra.get("raw"):
                    msg = {"role": "assistant", "content": answer}
                    if extra.get("raw"):
                        msg["anthropic_content"] = extra["raw"]  # 原样带回，Claude 需要它的思考块
                    history.append(msg)
                if not answer and stop != "refusal":
                    console.print(Text("  （AI 没有返回内容，可以再问一次）", style=DIM))
                if take_supplements(history):
                    continue  # 回答完了但用户又补充了：接着回答补充的内容
                break
            msg = {"role": "assistant", "content": answer, "tool_calls": calls}
            if extra.get("raw"):
                msg["anthropic_content"] = extra["raw"]
            elif reasoning:
                msg["reasoning_content"] = reasoning  # DeepSeek 思考模式调用工具时要求带回
            history.append(msg)
            # 每个工具调用都必须有对应的结果，否则下一次请求会被拒绝
            pending = list(calls)
            try:
                while pending:
                    call = pending.pop(0)
                    result, images = execute_tool(call, state)
                    msg = {"role": "tool", "tool_call_id": call["id"], "content": result}
                    if images:
                        msg["images"] = images
                    history.append(msg)
            except (KeyboardInterrupt, EOFError, Denied) as e:
                reason = "用户拒绝了这个操作，等待用户进一步指示。" if isinstance(e, Denied) else "用户中断了操作。"
                for c in [call] + pending:
                    if not any(m.get("tool_call_id") == c["id"] for m in history if m["role"] == "tool"):
                        history.append({"role": "tool", "tool_call_id": c["id"], "content": reason})
                if not isinstance(e, Denied):
                    console.print(Text("\n  （已停止）", style=DIM))
                elif take_supplements(history):
                    continue  # 拒绝之后用户补充了新的说法：按补充的接着做
                break
        else:
            console.print(Text(f"  已连续操作 {MAX_STEPS} 步，先停一下。可以说“继续”让它接着做。", style=DIM))
    except KeyboardInterrupt:
        pass
    finally:
        VIEW.stop()
    used = stats["tokens"] - tokens_before
    console.print(Text(f"  ⏱ {time.time() - started:.1f}s", style=DIM)
                  + (Text(f"  ·  {used} tokens", style=DIM) if used else Text()))
    console.print()
    return VIEW.typeahead.leftover()


# ───────────────────────── 会话记录 ─────────────────────────

SESSIONS_DIR = os.path.join(APP_DIR, "sessions")


def _session_path(sid):
    return os.path.join(SESSIONS_DIR, f"{sid}.json")


def save_session(sess, history, cfg):
    """每次回答完自动保存；sess["id"] 为空时新建一个会话。"""
    if not any(m["role"] == "user" for m in history):
        return
    os.makedirs(SESSIONS_DIR, exist_ok=True)
    now = time.time()
    if not sess.get("id"):
        sess["id"] = time.strftime("%Y%m%d-%H%M%S") + f"-{int(now * 1000) % 1000:03d}"
        sess["created"] = now
    first = next(m["content"] for m in history if m["role"] == "user")
    data = {
        "id": sess["id"],
        "title": re.sub(r"\s+", " ", first).strip()[:40],
        "created": sess.get("created", now),
        "updated": now,
        "model": cfg["model"],
        "profile": cfg.get("name"),
        "messages": history,
    }
    tmp = _session_path(sess["id"]) + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)
    os.chmod(tmp, 0o600)
    os.replace(tmp, _session_path(sess["id"]))


def list_sessions():
    items = []
    if not os.path.isdir(SESSIONS_DIR):
        return items
    for name in os.listdir(SESSIONS_DIR):
        if not name.endswith(".json"):
            continue
        try:
            with open(os.path.join(SESSIONS_DIR, name), encoding="utf-8") as f:
                item = json.load(f)
        except (OSError, json.JSONDecodeError):
            continue
        if isinstance(item, dict) and item.get("id") and isinstance(item.get("messages"), list):
            items.append(item)
    return sorted(items, key=lambda s: s.get("updated", 0), reverse=True)


def delete_session(sid):
    try:
        os.remove(_session_path(sid))
    except FileNotFoundError:
        pass


def _when(ts):
    t = time.localtime(ts)
    if time.strftime("%Y%m%d", t) == time.strftime("%Y%m%d"):
        return "今天 " + time.strftime("%H:%M", t)
    if time.strftime("%Y%m%d", t) == time.strftime("%Y%m%d", time.localtime(time.time() - 86400)):
        return "昨天 " + time.strftime("%H:%M", t)
    return time.strftime("%m-%d %H:%M", t)


def replay(history):
    """打开旧会话时，把聊天记录重新显示出来（工具调用只显示一行摘要）。"""
    results = {m.get("tool_call_id"): m.get("content", "") for m in history if m["role"] == "tool"}
    for m in history:
        if m["role"] == "user":
            console.print(Panel(Text(m["content"]), border_style=DIM, title="你", title_align="left"))
            for img in m.get("images") or []:
                console.print(Text(f"  🖼  {img.get('name', '图片')}", style=DIM))
        elif m["role"] == "assistant":
            if m.get("content"):
                console.print(Padding(md(m["content"]), (0, 0, 0, 2)))
            for call in m.get("tool_calls") or []:
                try:
                    args = json.loads(call["function"].get("arguments") or "{}")
                except json.JSONDecodeError:
                    args = {}
                if not isinstance(args, dict):
                    args = {}
                name = call["function"]["name"]
                what = args.get("command") or _short(_path(args.get("path")))
                icon = {"run_command": "⚡", "read_file": "📄", "list_dir": "📁"}.get(name, "✎")
                line = Text(f"  {icon} {what}", style=DIM)
                result = results.get(call.get("id"), "")
                if result.startswith("用户拒绝"):
                    line.append("  ✗ 已拒绝", style="red")
                elif result.startswith("用户中断"):
                    line.append("  （已中断）", style=DIM)
                console.print(line)
            if m.get("content"):
                console.print()


def history_menu(history, sess, cfg, stats):
    """/his：列出会话，可打开、删除一个、删除全部。返回 True 表示换了会话。"""
    while True:
        items = list_sessions()[:30]
        if not items:
            console.print(Text("  还没有保存的会话\n", style=DIM))
            return False
        table = Table(box=None, padding=(0, 2), show_header=False)
        for i, s in enumerate(items, 1):
            rounds = sum(1 for m in s.get("messages", []) if m["role"] == "user")
            mark = " ● 当前" if s["id"] == sess.get("id") else ""
            table.add_row(Text(f"{i}", style=f"bold {ACCENT}"), Text(s.get("title") or "（无标题）", style="bold"),
                          Text(f"{_when(s.get('updated', 0))} · {rounds} 轮{mark}", style=DIM))
        total = len(list_sessions())
        sub = f"共 {total} 个" + ("，显示最近 30 个" if total > 30 else "")
        console.print(Panel(table, title="历史会话", title_align="left", border_style=ACCENT,
                            subtitle=sub, subtitle_align="right"))
        ans = plain_prompt(HTML(
            "  <b>序号</b> 打开 · <b>d序号</b> 删除（如 d2）· <b>D</b> 删除全部 · 回车 返回 : ")).strip()
        if not ans:
            console.print()
            return False
        if ans.isdigit() and 1 <= int(ans) <= len(items):
            s = items[int(ans) - 1]
            history[:] = s.get("messages", [])
            sess.clear()
            sess.update(id=s["id"], created=s.get("created"))
            stats["tokens"] = 0
            show_banner(cfg)
            console.print(Text(f"  ↺ 已打开会话「{s.get('title')}」，可以接着聊\n", style="green"))
            replay(history)
            return True
        m = re.fullmatch(r"[dD]\s*(\d+)", ans)
        if m and 1 <= int(m.group(1)) <= len(items):
            s = items[int(m.group(1)) - 1]
            if plain_prompt(HTML(f"  确定删除「{s.get('title')}」吗？(y/n) : ")).strip().lower() in ("y", "yes", "是"):
                delete_session(s["id"])
                console.print(Text("  ✓ 已删除\n", style="green"))
                if s["id"] == sess.get("id"):
                    # 删的是正在聊的会话：一起清空，开始新对话
                    history.clear()
                    sess.clear()
                    stats["tokens"] = 0
                    show_banner(cfg)
                    console.print(Text("  ✓ 已删除当前会话，开始新对话\n", style="green"))
                    return True
            continue
        if ans == "D":
            total = len(list_sessions())
            confirm = plain_prompt(HTML(f"  <ansired>确定删除全部 {total} 个会话吗？删除后无法恢复。</ansired>"
                                        "输入 <b>yes</b> 确认 : ")).strip().lower()
            if confirm == "yes":
                for s in list_sessions():
                    delete_session(s["id"])
                history.clear()
                sess.clear()
                stats["tokens"] = 0
                show_banner(cfg)
                console.print(Text(f"  ✓ 已删除全部 {total} 个会话，开始新对话\n", style="green"))
                return True
            console.print(Text("  已取消\n", style=DIM))
            continue
        console.print(Text("  没看懂，请按提示输入", style="red"))


def memory_menu():
    """/memory：查看、添加、删除 AI 记住的事。"""
    while True:
        items = load_memory()
        if items:
            table = Table(box=None, padding=(0, 2), show_header=False)
            for i, m in enumerate(items, 1):
                table.add_row(Text(f"{i}", style=f"bold {ACCENT}"), Text(m["text"]), Text(m.get("time", ""), style=DIM))
        else:
            table = Text("还没有记住任何事。聊天时说“记住……”，AI 就会记下来。", style=DIM)
        console.print(Panel(table, title="🧠 AI 记住的事", title_align="left", border_style="magenta"))
        ans = plain_prompt(HTML("  <b>a 内容</b> 添加 · <b>d序号</b> 删除 · <b>D</b> 全部忘掉 · 回车 返回 : ")).strip()
        if not ans:
            console.print()
            return
        if ans[:2].lower() == "a ":
            console.print(Text(f"  ✓ {tool_remember(ans[2:])[0]}", style="green"))
            continue
        m = re.fullmatch(r"[dD]\s*(\d+)", ans)
        if m:
            text, ok = tool_forget_memory(m.group(1))
            console.print(Text(f"  {'✓' if ok else '✗'} {text}", style="green" if ok else "red"))
            continue
        if ans == "D" and items:
            if plain_prompt(HTML(f"  确定忘掉全部 {len(items)} 条吗？输入 <b>yes</b> 确认 : ")).strip().lower() == "yes":
                save_memory([])
                console.print(Text("  ✓ 已全部忘掉", style="green"))
            continue
        console.print(Text("  没看懂，请按提示输入", style="red"))


def copy_last_answer(history):
    answer = next((m["content"] for m in reversed(history) if m["role"] == "assistant" and m.get("content")), "")
    if not answer:
        console.print(Text("  还没有可以复制的回答\n", style=DIM))
        return
    subprocess.run(["pbcopy"], input=answer, text=True)
    console.print(Text(f"  ✓ 已复制上一条回答（{len(answer)} 字），可以去别的地方粘贴了\n", style="green"))


def pick_model(cfg):
    try:
        with console.status("正在获取模型列表…", spinner="dots"):
            models = list_models(cfg)
    except Exception as e:
        console.print(Panel(Text(f"获取模型列表失败：{e}"), title="出错了", border_style="red"), "\n")
        return
    default = models.index(cfg["model"]) + 1 if cfg["model"] in models else 1
    n = ask_choice("选择模型", [(m, "● 当前" if m == cfg["model"] else "") for m in models], default)
    cfg["model"] = models[n - 1]
    save_config(cfg)
    console.print(Text(f"  ✓ 已切换到 {cfg['model']}（下次打开也会用它）", style="green"), "\n")


def switch_line(cfg):
    if not cfg.get("backup_url"):
        console.print(Text("  没有设置备用线路，输入 /setup 可以换线路\n", style=DIM))
        return
    cfg["base_url"], cfg["backup_url"] = cfg["backup_url"], cfg["base_url"]
    save_config(cfg)
    console.print(Text(f"  ✓ 已切换到 {cfg['base_url']}\n", style="green"))


def main():
    args = sys.argv[1:]
    if args and args[0] in ("-v", "--version"):
        print(f"ds-chat {VERSION}")
        return
    if args and args[0] in ("--update", "update"):
        os.execvp("bash", ["bash", "-c", INSTALL_CMD])
    cfg = load_config()
    if cfg is None or (args and args[0] == "--setup"):
        try:
            new = setup_wizard(cfg)
        except (KeyboardInterrupt, EOFError):
            console.print(Text("\n  已取消设置\n", style=DIM))
            return
        if cfg is None:
            STORE.update(profiles=[new], active=0)
        else:
            STORE["profiles"][STORE["active"]] = new
        save_config()
        cfg = active_profile()
        if args and args[0] == "--setup":
            return

    history, stats = [], {"tokens": 0, "balance": None, "last_cost": None}
    state = {"auto_approve": False}  # 本次打开期间是否不再询问
    MENUBAR.start()
    sess = {}  # 当前会话的 id，第一次保存时生成
    refresh_balance(cfg, stats)

    def make_style():  # 每次输入前按当前主题色重新生成
        return Style.from_dict({
            "prompt": f"bold {ACCENT}",
            "box": ACCENT,
            "bottom-toolbar": "noreverse #808080 bg:default",
            "bottom-toolbar.box": f"noreverse {ACCENT} bg:default",
            "bottom-toolbar.accent": f"noreverse {ACCENT} bg:default",
            "bottom-toolbar.balance": "noreverse bold #4ade80 bg:default",
        })
    os.makedirs(APP_DIR, exist_ok=True)
    input_history = os.path.join(APP_DIR, "history")
    open(input_history, "a").close()
    os.chmod(input_history, 0o600)  # 里面是你输入过的所有内容
    session = PromptSession(history=FileHistory(input_history), style=make_style(),
                            cursor=CursorShape.BLINKING_BEAM, reserve_space_for_menu=0,
                            erase_when_done=True, refresh_interval=1)

    def width():
        return shutil.get_terminal_size().columns

    # 输入框：上边框单独打印，左边框在提示符里，右边框用 rprompt，下边框和状态栏放在 toolbar
    def box_left():
        return [("class:box", "│ "), ("class:prompt", "❯ ")]

    def toolbar():
        rounds = sum(1 for m in history if m["role"] == "user")
        left_accent = f"  ◆ {cfg['name']} · {cfg['model']}" + ("  ⚡Agent" if STORE.get("agent", True) else "")
        left = f"   {rounds} 轮对话   {stats['tokens']} tokens   /help 查看命令"
        bal = stats["balance"]
        cost = stats.get("last_cost")
        right = f"本次 💰{cost:.4f}   " if cost else ""
        if bal is None:
            right += "余额 查询中…  "
        elif isinstance(bal, float):
            right += f"💰{bal:.2f}  "
        # 按显示宽度（中文、⚡ 占两格）算右对齐的空格
        used = get_cwidth(left_accent + left + right)
        return [
            ("class:box", "╰" + "─" * (width() - 2) + "╯\n"),
            ("class:accent", left_accent),
            ("", left + " " * max(width() - used, 2)),
            ("class:balance", right),
        ]

    show_banner(cfg)
    prefill = ""  # AI 工作时补充了但没来得及发出去的内容，放回下一次的输入框
    if args and args[0] in ("-c", "--continue"):
        last = next(iter(list_sessions()), None)
        if last:
            history[:] = last.get("messages", [])
            sess.update(id=last["id"], created=last.get("created"))
            console.print(Text(f"  ↺ 接着上次的会话「{last.get('title')}」\n", style="green"))
            replay(history)
        else:
            console.print(Text("  还没有保存的会话，开始新对话\n", style=DIM))
    elif args and args[0] in ("-r", "--resume", "--history"):
        history_menu(history, sess, cfg, stats)
    while True:
        apply_theme(cfg)  # /api、/model、/setup 换了模型后，主题色跟着变
        console.print("╭" + "─" * (width() - 2) + "╮", style=ACCENT)
        try:
            text = session.prompt(box_left, rprompt=[("class:box", "│")], style=make_style(),
                                  bottom_toolbar=toolbar, default=prefill).strip()
            prefill = ""
        except KeyboardInterrupt:
            text = None
        except EOFError:
            break
        finally:
            # 擦掉上边框（输入框其余部分 prompt_toolkit 自己会擦掉）
            sys.stdout.write("\x1b[1A\x1b[2K")
            sys.stdout.flush()
        if not text:
            continue
        # 输入框提交后被擦掉，把用户的话换成一个灰色框留在记录里
        console.print(Panel(Text(text), border_style=DIM, title="你", title_align="left"))
        cmd = text.lower()
        try:
            if cmd in ("/exit", "/quit", "exit", "quit"):
                break
            if cmd in ("/his", "/history", "/h", "/resume", "/sessions"):
                if history_menu(history, sess, cfg, stats):
                    state["auto_approve"] = False
                continue
            if cmd in ("/clear", "/new"):
                history.clear()
                sess.clear()
                stats["tokens"] = 0
                state["auto_approve"] = False
                show_banner(cfg)
                continue
            if cmd in ("/model", "/m"):
                pick_model(cfg)
                continue
            if cmd == "/agent":
                STORE["agent"] = not STORE.get("agent", True)
                save_config()
                console.print(Text("  ✓ Agent 模式已打开：AI 可以运行命令、读写文件（改动前会问你）\n"
                                   if STORE["agent"] else "  ✓ 已切换为纯聊天模式\n", style="green"))
                continue
            if cmd in ("/memory", "/mem"):
                memory_menu()
                continue
            if cmd == "/copy":
                copy_last_answer(history)
                continue
            if cmd == "/compact":
                if len([m for m in history if m["role"] == "user"]) < 2:
                    console.print(Text("  对话还很短，不用压缩\n", style=DIM))
                    continue
                before_n = estimate_tokens(history)
                with console.status("正在压缩对话…", spinner="dots"):
                    ok = compact(cfg, history)
                console.print(Text(f"  ✓ 已压缩：约 {before_n} → {estimate_tokens(history)} tokens\n" if ok
                                   else "  压缩失败，稍后再试\n", style="green" if ok else "red"))
                if ok:
                    save_session(sess, history, cfg)
                continue
            if cmd == "/bar":
                STORE["menubar"] = not STORE.get("menubar", True)
                save_config()
                if STORE["menubar"]:
                    MENUBAR.start()
                    MENUBAR.update(stats["balance"])
                    shown = "已打开" if os.access(MenuBar.BIN, os.X_OK) else "已打开，但还没装好菜单栏小程序（重新运行 ai --update 安装）"
                else:
                    MENUBAR.stop()
                    shown = "已关闭"
                console.print(Text(f"  ✓ 菜单栏余额{shown}\n", style="green"))
                continue
            if cmd in ("/api", "/key", "/keys", "/apis"):
                old = cfg
                cfg = profile_menu(cfg)
                if cfg is not old or cfg.get("base_url") != old.get("base_url"):
                    stats.update(balance=None, last_cost=None)
                    refresh_balance(cfg, stats)
                continue
            if cmd == "/line":
                switch_line(cfg)
                refresh_balance(cfg, stats)
                continue
            if cmd == "/setup":
                STORE["profiles"][STORE["active"]] = setup_wizard(cfg)
                save_config()
                cfg = active_profile()
                stats["balance"] = None
                refresh_balance(cfg, stats)
                show_banner(cfg)
                if history:
                    replay(history)
                continue
            if cmd == "/help":
                show_help()
                continue
            if re.fullmatch(r"/[a-z]+", cmd):
                console.print(Text(f"  没有 {text} 这个命令，输入 /help 查看全部命令\n", style="red"))
                continue
        except (KeyboardInterrupt, EOFError):
            console.print()
            continue
        before = len(history)
        user_msg = {"role": "user", "content": text}
        images, img_errors = attach_images(text)
        for err in img_errors:
            console.print(Text(f"  ⚠ 图片读取失败：{err}", style="red"))
        if images:
            user_msg["images"] = images
            console.print(Text("  🖼  已附上图片：" + "、".join(i["name"] for i in images), style=DIM))
        history.append(user_msg)
        console.print()
        bal_before = stats["balance"]
        try:
            prefill = ask(cfg, history, stats, state) or ""
            refresh_balance(cfg, stats, before=bal_before)
        except urllib.error.HTTPError as e:
            del history[before:]
            detail = (getattr(e, "ds_body", None) or e.read().decode("utf-8", "replace"))[:300]
            tip = ""
            if images and e.code == 400:
                tip = "\n可能是当前模型不能看图片。用 /model 换名字里带 vision 的模型，或 /api 换成 Claude。"
            elif e.code >= 500 or e.code == 404:
                tip = "\n这个模型可能暂时不可用，输入 /model 换一个试试。"
            elif e.code in (401, 403):
                tip = "\n密钥无效或没有权限，输入 /setup 重新填写。"
            elif e.code == 429:
                tip = "\n请求太频繁或余额不足，稍等一下再试。"
            console.print(Panel(Text(f"HTTP {e.code}：{detail}{tip}"), title="出错了", border_style="red"), "\n")
        except (urllib.error.URLError, TimeoutError, OSError) as e:
            del history[before:]
            tip = "输入 /line 切换到备用线路再试试。" if cfg.get("backup_url") else "检查网络，或输入 /setup 换线路。"
            console.print(Panel(Text(f"网络连接失败：{e}\n{tip}"), title="出错了", border_style="red"), "\n")
        except Exception as e:  # 兜底：出什么错都不退出程序，对话也不丢（Claude SDK 的错误也在这里）
            del history[before:]
            code = error_code(e)
            tip = ""
            if images and code == 400:
                tip = "\n可能是当前模型不能看图片。用 /model 换名字里带 vision 的模型，或 /api 换成 Claude。"
            elif code in (401, 403):
                tip = "\n密钥无效或没有权限，输入 /setup 重新填写，或 /api 换一套配置。"
            elif code == 429:
                tip = "\n请求太频繁或额度不足，稍等一下再试。"
            elif code and code >= 500 or code == 404:
                tip = "\n这个模型可能暂时不可用，输入 /model 换一个，或 /api 换一套配置。"
            elif type(e).__name__ in ("APIConnectionError", "APITimeoutError"):
                tip = "\n网络连接失败，检查网络或代理后再试。"
            console.print(Panel(Text(f"{e}\n（{type(e).__name__}）{tip}"), title="出错了", border_style="red"), "\n")
        try:
            save_session(sess, history, cfg)
        except OSError as e:
            console.print(Text(f"  会话保存失败：{e}\n", style="red"))
    MENUBAR.stop()
    console.print(Text("\n  再见 👋\n", style=ACCENT))


if __name__ == "__main__":
    main()
